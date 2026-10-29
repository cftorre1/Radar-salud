"""Statistics fichas (Dirección 28-oct-2026): formulas, denominators, triggers, verifier, HOLD, GES.

No network and no model calls: the Lectura model is always a fake.
"""
import json
import shutil
from datetime import date
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest

from radar_salud import data_insights, stats_ficha_f1 as f1, stats_fichas as sf

REPO = Path(__file__).resolve().parents[1]
TODAY = date(2026, 10, 28)
INPUTS = ("data/market/isapre_portfolio.json", "data/excel/validated_series.json", "data/entities/isapres.json",
          "data/history/superintendencia_signals.json")


@pytest.fixture(autouse=True)
def isolated_ledgers(tmp_path, monkeypatch):
    """Budget counters (ai_budget) write under RADAR_ROOT: never touch the real data/ai_usage ledger."""
    monkeypatch.setenv("RADAR_ROOT", str(tmp_path / "ledger-root"))


@pytest.fixture
def root(tmp_path):
    for rel in INPUTS:
        (tmp_path / rel).parent.mkdir(parents=True, exist_ok=True)
        shutil.copy(REPO / rel, tmp_path / rel)
    return tmp_path


def portfolio(root):
    return json.loads((root / INPUTS[0]).read_text(encoding="utf-8"))


def save_portfolio(root, data):
    (root / INPUTS[0]).write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")


def fake_model(text, cost=0.0004, review="PASS", claim=""):
    """Writer and reviewer fake: the reviewer call (feature *_review) answers ``review`` with ``claim``."""
    calls = []

    def call(prompt, **kw):
        calls.append({"prompt": prompt, **kw})
        if str(kw.get("feature", "")).endswith("_review"):
            return SimpleNamespace(output={"verdict": review, "claim": claim}, model="gpt-5.6-luna",
                                   estimated_cost_usd=cost / 4, response_id="rv1")
        return SimpleNamespace(output={"lectura": text}, model="gpt-5.6-luna", estimated_cost_usd=cost, response_id="r1")
    call.calls = calls
    call.writes = lambda: [c for c in calls if not c["feature"].endswith("_review")]
    call.reviews = lambda: [c for c in calls if c["feature"].endswith("_review")]
    return call


ALLOW = lambda worst: (True, "ok")  # noqa: E731
GOOD = ("Para la base de ingresos importan más los cotizantes que los beneficiarios: quien paga se mantiene y lo que "
        "se reduce son cargas. Para la competencia, la pregunta es si Banmédica estabiliza su cartera antes de que "
        "Esencial y Nueva Masvida consoliden lo ganado.")


# ---------------------------------------------------------------------------------------------
# Formulas and denominators


def test_package_formulas_use_stocks_open_denominator_and_reconcile(root):
    p = f1.build_package(root)
    v = p["values"]
    assert p["reconciliation"]["ok"] is True and p["period"] == "2026-08"
    assert v["sistema.beneficiarios.2026-08"]["value"] == 2486683 == v["cartera.declarado.beneficiarios.2026-08"]["value"]
    assert v["sistema.beneficiarios.mm.2026-08"]["value"] == -814
    # Stock rule: the year change is end minus start, never a sum of months.
    assert v["sistema.beneficiarios.ytd"]["value"] == 2486683 - 2514038 == -27355
    assert v["sistema.cotizantes.mm.2026-08"]["value"] == 1296 and v["sistema.cotizantes.mm.2026-07"]["value"] == 475
    # Market share: only the 7 open Isapres in the denominator.
    assert v["abiertas.beneficiarios.2026-08"]["value"] == 2419986
    assert len(v["abiertas.beneficiarios.2026-08"]["inputs"]) == 7
    share_sum = sum(v[f"share.{c}.2026-08"]["value"] for c in (67, 78, 80, 81, 99, 107, 108))
    assert abs(share_sum - 100) < 1e-9
    assert "share.63.2026-08" not in v  # closed Isapres never get a market share
    assert round(v["contrib.99"]["value"]) == 77 and round(v["cargas.contrib"]["value"]) == 78
    assert round(v["mezcla.2026-08"]["value"], 3) == 0.603 and round(v["mezcla.2026-01"]["value"], 3) == 0.615
    assert v["movilidad.diferencia_intervalo"]["value"] == -17700
    for item in v.values():  # provenance on every captured value
        if item["kind"] == "captured":
            assert item["locator"] and item["source"] in p["sources"]
    assert all(sf.HEX64.fullmatch(s["sha256"]) for s in p["sources"].values())


def test_contribution_not_computed_when_system_change_is_small(root):
    data = portfolio(root)
    # Make Aug equal Jan for every code: |Δ system| = 0 < 5.000.
    jan = {(o["official_code"], o["metric_id"]): o["value"] for o in data["observations"] if o["period"] == "2026-01"}
    for o in data["observations"]:
        if o["period"] == "2026-08" and o["metric_id"] in ("beneficiarios", "cotizantes", "cargas"):
            o["value"] = jan[(o["official_code"], o["metric_id"])]
    data["coverage"][-1]["total_beneficiaries"] = data["coverage"][0]["total_beneficiaries"]
    save_portfolio(root, data)
    p = f1.build_package(root)
    assert p["reconciliation"]["ok"] is True
    assert not any(k.startswith("contrib.") for k in p["values"])
    assert not any(f["type"] == "concentration" for f in p["findings"])


def test_ficha_shows_august_and_keeps_yoy_pending(root):
    r = f1.run(root, today=TODAY, llm=False)
    assert r["status"] == "pass"
    fi = r["ficha"]
    assert fi["period"] == "2026-08" and "agosto 2026" in fi["kicker_meta"]
    assert fi["headline_type"] == "record_low_decline"
    # Stock wording: a change between month-end stocks is not a flow of exits (those live in the mobility tab).
    assert fi["headline"] == "La cartera Isapre casi deja de caer: −814 beneficiarios en agosto, la menor baja del año"
    assert not any(w in fi["headline"].lower() for w in ("salida", "afiliados", "se detiene"))
    assert "2.486.683" in fi["dek"] and len(fi["dek"].split()) <= 40
    assert fi["kpis"][0]["badge"] == "Sin umbral cruzado · a/a pendiente"
    assert any("pendiente" in n["text"] for n in fi["notes"])
    tabs = {t["id"]: t for t in fi["tabs"]}
    assert list(tabs) == ["total", "isapres", "mezcla", "movilidad", "regional"]
    # Dirección 29-oct (C→B): stock evolution first, monthly differences second.
    assert [c["form"] for c in tabs["total"]["charts"]] == ["line", "columns_signed"]
    assert [row["label"] for row in tabs["total"]["charts"][0]["rows"]][-1] == "Ago"
    share = tabs["isapres"]["charts"][1]
    assert share["form"] == "dots" and len(share["rows"]) == 7  # open Isapres only
    assert tabs["regional"]["status"] == "pendiente" and tabs["regional"]["links"][0]["url"].endswith("a-nivel-regional-agosto-2026/")
    assert fi["verification"]["verdict"] == "PASS" and fi["featured"] is True
    assert fi["verification"]["lenses"] == {"data": "PASS", "editorial": "PASS", "ui": "PASS", "business": "PASS"}
    assert fi["lectura"]["status"] == "not_attempted"
    # Facts stay as data for the Lectura; the topic layout shows text only where it adds something (Dirección 29-oct).
    assert fi["facts"]["rendered_in_tabs"] is True and len(fi["facts"]["paragraphs"]) == 3
    assert len(fi["watch"]) == 2


def test_topics_follow_the_29_oct_content_rules(root):
    """Index → topic (option C→B): index line + sparkline per topic; dek and KPI cards only where they add something."""
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    tabs = {t["id"]: t for t in fi["tabs"]}
    assert {k: t["index_line"] for k, t in tabs.items()} == {
        "total": "2.486.683 beneficiarios · menor baja del año", "isapres": "Banmédica: 77% de la baja · 3 Isapres crecen",
        "mezcla": "Cotizantes suben por segundo mes", "movilidad": "Salen más de los que entran · dato a julio",
        "regional": "Enlace a la SIS · cifras aún no extraídas"}
    assert {k: t["spark"] for k, t in tabs.items()} == {
        "total": {"chart": "f1-stock", "form": "line"}, "isapres": {"chart": "f1-isapres", "form": "diverging"},
        "mezcla": {"chart": "f1-cot", "form": "columns"}, "movilidad": {"chart": "f1-mob", "form": "diverging"}, "regional": None}
    for k in ("total", "isapres", "movilidad"):
        assert tabs[k]["dek"] is None and tabs[k]["kpis"] == [] and tabs[k]["text"] == []
    assert tabs["mezcla"]["dek"] == "Las cargas explican el 78% de la baja del año."
    assert [(c["label"], c["value"], c["lines"]) for c in tabs["mezcla"]["kpis"]] == [
        ("Cotizantes", "1.551.169", ["−5.895 desde enero (−0,38%)"]), ("Cargas", "935.514", ["−21.460 desde enero (−2,24%)"])]
    assert "dek" not in tabs["regional"] and tabs["regional"]["status"] == "pendiente"
    # Chart text: reading instructions and method clarifications removed; keys and provenance kept.
    charts = {c["id"]: c for t in fi["tabs"] for c in t["charts"]}
    texts = " ".join(str(c.get(k) or "") for c in charts.values() for k in ("headline", "subtitle", "legend", "source")).lower()
    for gone in ("eje sin partir de cero", "columna más corta", "tinta = pierde", "agosto en acento", "no se suman"):
        assert gone not in texts, gone
    assert not any(c.get("legend") for c in charts.values())
    assert "punto = agosto, aro = enero" in charts["f1-share"]["subtitle"]
    assert "la calcula alicanto" in charts["f1-mm"]["source"].lower()
    assert charts["f1-stock"]["headline"] == "Beneficiarios Isapre, enero a agosto"
    # The by-Isapre table: Cambio = Cotizantes + Cargas on every row, closed Isapres grouped, a Sistema row.
    tbl = charts["f1-isapres"]["table"]
    assert tbl["columns"] == ["Isapre", "Cambio", "Cotizantes", "Cargas", "% vs. enero"]
    num = lambda x: int(x.replace("−", "-").replace("+", "").replace(".", ""))  # noqa: E731
    for row in tbl["rows"]:
        assert num(row[1]) == num(row[2]) + num(row[3]), row
    assert [r[0] for r in tbl["rows"]][-1] == "Sistema" and "Cerradas (3)" in [r[0] for r in tbl["rows"]]
    assert tbl["rows"][0] == ["Banmédica", "−20.948", "−9.542", "−11.406", "−3,70%"]
    assert tbl["rows"][-1] == ["Sistema", "−27.355", "−5.895", "−21.460", "−1,09%"]
    assert any("'Menor baja del año' compara" in n["text"] for n in fi["notes"])


def test_verify_ficha_traces_topic_kpis_index_lines_and_sparks(root):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    p = f1.build_package(root)
    assert sf.verify_ficha(fi, p)["verdict"] == "PASS"

    def broken(edit):
        tabs = json.loads(json.dumps(fi["tabs"]))
        edit({t["id"]: t for t in tabs})
        return sf.verify_ficha({**fi, "tabs": tabs}, p)["errors"]

    # A figure on a topic card must come from the values that card cites, not from any package value.
    assert any(e.startswith("topic_kpi_number_not_in_refs:mezcla:tema-cargas")
               for e in broken(lambda t: t["mezcla"]["kpis"][1].update(value="1.551.169")))
    assert any(e.startswith("topic_kpi_ref_missing:mezcla") for e in broken(lambda t: t["mezcla"]["kpis"][0].update(refs=["no.existe"])))
    assert "number_not_in_evidence:81" in broken(lambda t: t["mezcla"].update(dek="Las cargas explican el 81% de la baja del año."))
    assert "count_not_in_evidence:isapres:4 crecen" in broken(lambda t: t["isapres"].update(index_line="Banmédica: 77% de la baja · 4 Isapres crecen"))
    assert "spark_not_a_topic_chart:total" in broken(lambda t: t["total"].update(spark={"chart": "f1-cot", "form": "line"}))
    assert "ui:index_line_long:total" in broken(lambda t: t["total"].update(index_line="x" * 70))
    assert "ui:topic_over_2_kpis:mezcla" in broken(lambda t: t["mezcla"]["kpis"].append(t["mezcla"]["kpis"][0]))


def test_topic_lines_cite_their_own_refs(root):
    """Independent review 29-oct: index lines and topic deks trace to the refs they declare (unit- and entity-aware)."""
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    p = f1.build_package(root)
    tabs = {t["id"]: t for t in fi["tabs"]}
    assert tabs["total"]["index_refs"] == ["sistema.beneficiarios.2026-08"]
    assert tabs["isapres"]["index_refs"] == ["contrib.99", "conteo.abiertas_crecen"]
    assert tabs["mezcla"]["index_refs"] == [] and tabs["mezcla"]["dek_refs"] == ["cargas.contrib"]

    def errors(edit):
        tabs2 = json.loads(json.dumps(fi["tabs"]))
        edit({t["id"]: t for t in tabs2})
        return sf.verify_ficha({**fi, "tabs": tabs2}, p)["errors"]

    # Some package value equals 79 (Isapre 76, cotizantes −79): a global pool would let this wrong share through.
    assert any(abs(v["value"]) == 79 for v in p["values"].values())
    assert "topic_number_not_in_refs:isapres:index:79%" in errors(
        lambda t: t["isapres"].update(index_line="Banmédica: 79% de la baja · 3 Isapres crecen"))
    # '15 regiones' must not trace to a share (or to anything) the line does not cite.
    assert "topic_number_not_in_refs:regional:index:15" in errors(lambda t: t["regional"].update(index_line="15 regiones con datos"))
    # Unit-aware: the contribution is a %, so a plain '77' does not match it, and '3' never matches a %.
    assert "topic_number_not_in_refs:isapres:index:77" in errors(lambda t: t["isapres"].update(index_line="Banmédica: 77 de la baja"))
    assert "topic_number_not_in_refs:isapres:index:3%" in errors(lambda t: t["isapres"].update(index_line="Banmédica: 3% de la baja"))
    # Entity-aware: Banmédica's share credited to another Isapre fails.
    other = next(e["name"] for e in p["entities"] if e["code"] == 67)
    assert f"topic_entity_not_in_refs:isapres:index:{other}" in errors(
        lambda t: t["isapres"].update(index_line=f"{other}: 77% de la baja"))
    # A dek figure traces to the dek's refs only.
    assert "topic_number_not_in_refs:mezcla:dek:81%" in errors(lambda t: t["mezcla"].update(dek="Las cargas explican el 81% de la baja del año."))
    assert any(e.startswith("topic_ref_missing:total:index") for e in errors(lambda t: t["total"].update(index_refs=["no.existe"])))


@pytest.mark.parametrize("tab,line,error", [
    ("isapres", "Banmédica: 79% de la baja · 3 Isapres crecen", "topic_number_not_in_refs:isapres:index:79%"),
    ("regional", "15 regiones con datos", "topic_number_not_in_refs:regional:index:15"),
    ("mezcla", "Cotizantes suben por tercer mes", "derived_claim_not_in_series:racha:cotizantes:tercer"),
    ("mezcla", "Cotizantes bajan por segundo mes", "derived_claim_not_in_series:racha:cotizantes:segundo"),
    ("total", "2.486.683 beneficiarios · menor baja del año", None),
])
def test_wrong_topic_lines_hold(root, monkeypatch, tab, line, error):
    """End to end: a wrong index line never publishes (HOLD); the true line still passes."""
    real = f1.compose

    def compose(*a, **kw):
        fi = real(*a, **kw)
        next(t for t in fi["tabs"] if t["id"] == tab)["index_line"] = line
        return fi
    monkeypatch.setattr(f1, "compose", compose)
    r = f1.run(root, today=TODAY, llm=False)
    if error is None:
        assert r["status"] == "pass", r.get("errors")
    else:
        assert r["status"] == "hold" and error in r["errors"], r.get("errors")


def test_derived_claims_are_recomputed_from_the_series(root):
    p = f1.build_package(root)
    assert sf.derived_claim_errors("Cotizantes suben por segundo mes", p) == []
    assert sf.derived_claim_errors("Segundo mes al alza (jul +475)", p, "cotizantes") == []
    assert sf.derived_claim_errors("la menor baja del año", p) == []
    assert sf.derived_claim_errors("Cotizantes: menor baja del año", p) == ["derived_claim_not_in_series:menor_baja:cotizantes"]
    worse = json.loads(json.dumps(p))
    worse["values"]["sistema.beneficiarios.mm.2026-08"]["value"] = -99999  # August is no longer the smallest decline
    assert sf.derived_claim_errors("la menor baja del año", worse) == ["derived_claim_not_in_series:menor_baja:beneficiarios"]
    late = {**p, "periods": p["periods"][1:]}  # series starts after January: 'del año' cannot be claimed
    assert "derived_claim_not_in_series:menor_baja:beneficiarios" in sf.derived_claim_errors("menor baja del año", late)
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    wrong = json.loads(json.dumps(fi))
    wrong["kpis"][1]["lines"][1] = "Tercer mes al alza (jul +475)"
    assert "derived_claim_not_in_series:racha:cotizantes:tercer" in sf.verify_ficha(wrong, p)["errors"]


def test_isapre_table_split_must_add_up(root):
    """A capture whose cotizantes + cargas do not add up to beneficiarios never publishes the split table."""
    data = portfolio(root)
    for o in data["observations"]:
        if o["period"] == "2026-08" and o["official_code"] == 99 and o["metric_id"] == "cargas":
            o["value"] += 1
    (root / INPUTS[0]).write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    assert f1.run(root, today=TODAY, llm=False)["status"] == "hold"


def test_verify_ficha_rejects_invented_sources_links_and_placeholders(root):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    p = f1.build_package(root)
    invented = {**fi, "sources": fi["sources"] + [{"label": "Inventada", "url": "https://example.cl/x", "sha256": "e" * 64}]}
    assert "source_not_captured:Inventada" in sf.verify_ficha(invented, p)["errors"]
    tabs = json.loads(json.dumps(fi["tabs"]))
    tabs[-1]["links"] = [{"label": "Otra", "url": "https://example.cl/regional.xlsx"}]
    assert "link_not_captured:regional" in sf.verify_ficha({**fi, "tabs": tabs}, p)["errors"]
    placeholder = {**fi, "dek": fi["dek"] + " {entity} ejemplo ilustrativo."}
    assert any(e.startswith("template_placeholder") for e in sf.verify_ficha(placeholder, p)["errors"])
    assert sf.verify_ficha(fi, p)["verdict"] == "PASS"


def test_missing_entity_name_holds_instead_of_showing_a_code(root):
    ents = json.loads((root / INPUTS[2]).read_text(encoding="utf-8"))
    for e in ents["entities"]:
        if e["official_code"] == 108:
            e["commercial_name"] = None
            e.pop("legal_name", None)
    (root / INPUTS[2]).write_text(json.dumps(ents, ensure_ascii=False), encoding="utf-8")
    r = f1.run(root, today=TODAY, llm=False)
    assert r["status"] == "hold" and "entity_name_missing" in r["errors"][0]


# ---------------------------------------------------------------------------------------------
# Triggers and HOLD


def test_triggers_new_period_then_no_change_then_revision(root):
    first = f1.run(root, today=TODAY, llm=False)
    sf.publish(first, root)
    again = f1.run(root, today=TODAY, llm=False)
    assert again["status"] == "no_change"
    # A new release that revises July by more than 0,5 % republishes with a 'revisado' note.
    data = portfolio(root)
    row = next(o for o in data["observations"] if o["period"] == "2026-07" and o["official_code"] == 99 and o["metric_id"] == "cargas")
    benef = next(o for o in data["observations"] if o["period"] == "2026-07" and o["official_code"] == 99 and o["metric_id"] == "beneficiarios")
    row["value"] += 20000
    benef["value"] += 20000
    data["coverage"][-2]["total_beneficiaries"] += 20000
    for o in data["observations"]:
        o["source_sha256"] = "b" * 64
    save_portfolio(root, data)
    revised = f1.run(root, today=TODAY, llm=False)
    assert revised["status"] == "pass"
    assert "revision" in revised["decision"]["triggers"]
    assert any(n["text"].startswith("Revisado (Total Sistema, beneficiarios, julio 2026): 2.487.497 → 2.507.497") for n in revised["ficha"]["notes"])


def test_hold_on_reconciliation_failure_keeps_live(root):
    sf.publish(f1.run(root, today=TODAY, llm=False), root)
    live_before = sf.load_live(root, "f1_cartera")
    data = portfolio(root)
    data["coverage"][-1]["total_beneficiaries"] += 1  # Σ by Isapre ≠ declared total
    for o in data["observations"]:
        o["source_sha256"] = "c" * 64
    save_portfolio(root, data)
    r = f1.run(root, today=TODAY, llm=False)
    assert r["status"] == "hold" and any(e.startswith("HOLD_RECONCILIATION") and "total_2026-08" in e for e in r["errors"])
    out = sf.publish(r, root)
    assert out["published"] is False and Path(out["hold"]).exists()
    assert sf.load_live(root, "f1_cartera") == live_before


def test_hold_on_missing_hash_and_older_period():
    base = {"reconciliation": {"ok": True}, "primary_source": {"sha256": "a" * 64}, "period": "2026-08",
            "release_key": "2026-08:x", "published_at": "2026-10-02", "findings": [], "revisions": []}
    cal = {"expected_release_interval_days": 31, "late_tolerance_days": 21}
    assert sf.evaluate_triggers({**base, "primary_source": {"sha256": ""}}, None, today=TODAY, calendar=cal)["holds"] == ["HOLD_MISSING_HASH"]
    older = sf.evaluate_triggers({**base, "period": "2026-07"}, {"period": "2026-08", "release_key": "z"}, today=TODAY, calendar=cal)
    assert older["decision"] == "HOLD" and "HOLD_OLDER_PERIOD" in older["holds"]
    fresh = sf.evaluate_triggers(base, None, today=TODAY, calendar=cal)
    assert fresh["decision"] == "PUBLISH" and fresh["triggers"] == ["new_period"] and fresh["featured"] is False


def test_late_release_adds_watch_item_and_republishes_once(root):
    sf.publish(f1.run(root, today=TODAY, llm=False), root)
    late_day = date(2026, 11, 30)  # 2 oct + 31 + 21 = 23 nov
    r = f1.run(root, today=late_day, llm=False)
    assert r["status"] == "pass" and "late_release" in r["decision"]["triggers"]
    assert "atrasada" in r["ficha"]["watch"][0]["headline"]
    sf.publish(r, root)
    assert f1.run(root, today=late_day, llm=False)["status"] == "no_change"


def test_publish_archives_previous_live(root):
    sf.publish(f1.run(root, today=TODAY, llm=False), root)
    data = portfolio(root)
    for o in data["observations"]:
        o["source_sha256"] = "d" * 64
    save_portfolio(root, data)
    out = sf.publish(f1.run(root, today=TODAY, llm=False), root)
    assert out["published"] is True and out["archived"] and Path(out["archived"]).exists()


# ---------------------------------------------------------------------------------------------
# Lectura verifier and budgeted generation (fake model only)


def test_verifier_accepts_clean_text_and_rejects_violations(root):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    p = f1.build_package(root)
    assert sf.verify_lectura(GOOD, p, ficha=fi)["verdict"] == "PASS"
    causal = sf.verify_lectura("La baja se debe a la ley corta y el alza de precios provoca salidas.", p, ficha=fi)
    assert any(e.startswith("causal_without_note") for e in causal["errors"])
    assert sf.verify_lectura("La lectura general de la cartera es estable.", p, ficha=fi)["verdict"] == "PASS"  # 'general' is not 'genera'
    other = sf.verify_lectura("Banmédica pierde cartera.", {**p, "entity_codes": [67]}, ficha=fi)
    assert "insurer_not_in_package:banmedica" in other["errors"]
    assert any(e.startswith("over_60_words") for e in sf.verify_lectura("cartera " * 61, p, ficha=fi)["errors"])
    with_note = sf.verify_lectura("La baja se debe a cargas.", {**p, "causal_note": "Nota causal verificada"}, ficha=fi)
    assert not any(e.startswith("causal") for e in with_note["errors"])
    assert "url_in_lectura" in sf.verify_lectura("Para la cartera ver www.example.cl", p, ficha=fi)["errors"]


# Figure rule (Dirección 28-oct 15:25): figures are allowed if they trace and add something.
# Every probe of the second independent review, positive and negative.
@pytest.mark.parametrize("text,error", [
    # traced and new (Banmédica share change, not shown in the ficha)
    ("Para la competencia, Banmédica cedió 0,6 pp de participación entre las abiertas desde enero.", None),
    ("Para la competencia, Banmédica cedió −0,6 pp de participación entre las abiertas desde enero.", None),
    ("Para la competencia, Esencial es la que más crece: 13,3% desde enero.", None),
    ("Para la competencia, Nueva Masvida gana cuatro décimas de participación entre las abiertas desde enero.", None),
    ("Para las Isapres, por cada entrada al sistema hubo 1,19 salidas entre los cortes.", None),  # derived ratio
    ("Para las Isapres, las cerradas perdieron 1.433 beneficiarios desde enero.", None),
    # wrong sign, unit, entity, period
    ("Para la competencia, Banmédica ganó 0,6 pp de participación entre las abiertas desde enero.", "number_sign_mismatch"),
    ("Para la competencia, Banmédica cedió +0,6 pp de participación entre las abiertas.", "number_sign_mismatch"),
    ("La cartera anotó +27.355 beneficiarios desde enero para las Isapres.", "number_sign_mismatch"),
    ("Para la competencia, Banmédica cedió 0,6% de participación entre las abiertas desde enero.", "number_unit_mismatch"),
    ("Para la competencia, Banmédica perdió 20.948 cotizantes desde enero.", "number_unit_mismatch"),
    ("Para la competencia, Esencial cedió 0,6 pp de participación entre las abiertas desde enero.", "number_entity_mismatch"),
    ("Para la competencia, Banmédica tenía 22,5% de las abiertas en enero.", "number_period_mismatch"),
    # restated figures (already in headline, dek, KPIs, tabs or facts)
    ("Banmédica concentra el 77% de la baja neta de la cartera desde enero.", "restates_shown"),
    ("La cartera perdió 814 beneficiarios.", "restates_shown"),
    ("Para las Isapres, tres de cada cuatro beneficiarios perdidos desde enero son de Banmédica.", "restates_shown"),  # 77% in words
    ("Para la competencia, Banmédica perdió miles de beneficiarios desde enero.", "restates_shown"),
    # untraced figures and number words
    ("La cartera cayó 3% en el año.", "number_"),  # 3 % is only an Isapre figure, not the system
    ("Veinte mil beneficiarios dejaron la cartera de las Isapres.", "number_not_in_pool"),
    ("La cartera perdió 814 personas en 2027.", "year_not_in_package"),
    ("Para inversionistas, un par de Isapres abiertas sostiene la cartera.", "number_unit_mismatch"),  # 2 is a rank, not a count
    ("Para las aseguradoras, el IV trimestre dirá si la cartera se estabiliza.", "period_not_in_package"),
    ("Para las aseguradoras, la cartera del siglo XXI cambia.", "number_not_in_pool"),
    ("Para la competencia, la baja de Banmédica es quíntuple que la de Esencial.", "number_not_in_pool"),
    ("Para la competencia, Banmédica ocupa el centésimo lugar en cartera.", "number_not_in_pool"),
    ("Para la competencia, Banmédica cede medio punto de participación entre las abiertas desde enero.", "number_not_in_pool"),
    ("Para la competencia, Nueva Masvida gana una décima de participación entre las abiertas desde enero.", "number_"),
    ("Para la competencia, uno de cada cuatro beneficiarios perdidos por la cartera son cargas.", "number_not_in_pool"),
    ("Para la competencia, la baja de Banmédica es el doble que la de Esencial.", "number_not_in_pool"),
    ("La cartera de las Isapres cae ½ punto.", "number_"),
    ("La cartera de las Isapres cae ٣ puntos.", "number_"),
    ("La cartera de las Isapres baja por segundo mes.", "restates_shown"),  # 'segundo mes' is the KPI
])
def test_lectura_figure_probes(root, text, error):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    r = sf.verify_lectura(text, f1.build_package(root), ficha=fi)
    if error is None:
        assert r["verdict"] == "PASS", r
    else:
        assert r["verdict"] == "FAIL" and any(e.startswith(error) for e in r["errors"]), r


# Third independent review (28-oct): every probe, negative and positive, as a class of failure.
@pytest.mark.parametrize("text,error", [
    # 1. scope: a system / open figure is never credited to an insurer
    ("Para las Isapres, Esencial tiene 935.514 cargas en agosto.", "number_entity_mismatch"),
    ("Para las Isapres, Banmédica tiene 2.419.986 beneficiarios.", "number_entity_mismatch"),
    ("Para las Isapres, las cerradas tienen 2.419.986 beneficiarios.", "number_entity_mismatch"),
    ("Para la competencia, Banmédica perdió 27.355 beneficiarios desde enero.", "number_entity_mismatch"),
    # 2. period kind: a month named matches only that month's change or stock; YTD needs YTD wording
    ("Para las Isapres, en agosto las Isapres cerradas perdieron 1.433 beneficiarios.", "number_period_mismatch"),
    ("Para las Isapres, las cerradas perdieron 1.433 beneficiarios.", "number_period_mismatch"),
    ("Para las Isapres, en julio el sistema perdió 814 beneficiarios.", "number_period_mismatch"),
    ("Para las Isapres, desde enero el sistema perdió 814 beneficiarios.", "number_period_mismatch"),
    ("Para las Isapres, en agosto salieron 109.663 cotizantes del sistema.", "number_period_mismatch"),
    # 3. comparators against the true value
    ("Para las Isapres, el sistema tiene más de 2,5 millones de beneficiarios.", "number_"),
    ("Para las Isapres, las abiertas tienen menos de 2,4 millones de beneficiarios.", "number_"),
    ("Para las Isapres, las abiertas tienen casi 2,4 millones de beneficiarios.", "number_"),
    ("Para las Isapres, el sistema tiene más de 1 millón de beneficiarios.", "number_"),  # true, but no information
    ("Para las Isapres, las abiertas tienen cerca de 2,6 millones de beneficiarios.", "number_"),
    # 4. gaming 'adds something': trivial counts and restatements with a new id
    ("Banmédica concentra el 77% de la baja neta, entre 7 Isapres abiertas.", "number_"),
    ("Banmédica concentra el 77% de la baja neta desde enero, entre 7 Isapres abiertas.", "restates_shown"),
    ("Para la competencia, Banmédica concentra el 77% de la baja desde enero y tres Isapres abiertas crecen.", "restates_shown"),
    ("La cartera Isapre casi deja de caer: −814 beneficiarios en agosto, la menor baja del año, con 2.419.986 "
     "beneficiarios en las abiertas.", "restates_shown_text"),
    # 5. insurers: misspellings are insurers; Fonasa (any alias) is outside the package
    ("Para las Isapres, el seguro público gana espacio frente a la cartera.", "insurer_not_in_package:fonasa"),
    ("Para las Isapres, la aseguradora estatal gana espacio frente a la cartera.", "insurer_not_in_package:fonasa"),
    ("Para la competencia, CB mantiene su cartera.", "insurer_not_in_facts:cruz_blanca"),
    ("Para la competencia, Consalut mantiene su cartera.", "insurer_not_in_facts:consalud"),
    ("Para la competencia, Colmenna mantiene su cartera.", "insurer_not_in_facts:colmena"),
    ("Para la competencia, Banmédika tiene 2.419.986 beneficiarios.", "number_entity_mismatch"),
    # 6. bare domains
    ("Para la cartera, ver superdesalud.gob.cl.", "url_in_lectura"),
    ("Para la cartera, ver example.com/informe.", "url_in_lectura"),
    # 8. ordinals: right ranking, right period
    ("Para la competencia, Nueva Masvida es la segunda Isapre con mayor baja desde enero.", "number_"),
    ("Para la competencia, la segunda Isapre con mayor baja pierde cartera en agosto.", "number_period_mismatch"),
])
def test_third_review_probes_fail(root, text, error):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    r = sf.verify_lectura(text, f1.build_package(root), ficha=fi)
    assert r["verdict"] == "FAIL" and any(e.startswith(error) for e in r["errors"]), r


# The reviewer's good sentences and three good Lecturas: figures that trace, add something and pass.
GOOD_LECTURAS = [
    "Para el mercado, los cotizantes subieron en 1.296 en agosto y el sistema sumó 1,3 contratos por cada desahucio "
    "voluntario en julio: la base que paga cotización resiste mejor que las cargas.",
    "Para la competencia, Banmédica pierde 2,4 veces lo que la segunda Isapre con mayor baja desde enero, y Esencial "
    "gana 1,28 veces lo de Nueva Masvida.",
    "Para inversionistas, las cerradas cayeron 2,1% desde enero, más que el 1,09% del sistema: el ajuste no se limita "
    "a las abiertas.",
]


@pytest.mark.parametrize("text", GOOD_LECTURAS + [
    "Para las Isapres, las abiertas tienen más de 2,4 millones de beneficiarios.",  # true 2.419.986
    "Para las Isapres, las abiertas tienen menos de 2,5 millones de beneficiarios.",  # true 2.419.986
    "Para las Isapres, las abiertas tienen cerca de 2,4 millones de beneficiarios.",
    "Para el mercado, los cotizantes subieron en 1.296 en agosto, mientras las cerradas perdieron 1.433 desde enero.",
    "Para las Isapres, el sistema anotó 1,3 contratos por cada desahucio voluntario en julio.",
    "Para el mercado, Esencial crece 13,3% desde enero.",
    "Para la competencia, Colmena es la segunda Isapre con mayor baja desde enero.".replace("Colmena", "Banmédica").replace(
        "segunda", "primera"),
    "Para las Isapres, por cada entrada al sistema hubo 1,19 salidas entre julio 2025 y julio 2026.",
    "Para la competencia, Escencial gana cartera frente a Banmédica.",  # misspelling = Esencial, which is in the facts
])
def test_third_review_good_sentences_pass(root, text):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    r = sf.verify_lectura(text, f1.build_package(root), ficha=fi)
    assert r["verdict"] == "PASS", r


# Fourth independent review (28-oct): the share.<code>.* figures are shares among the OPEN Isapres only
# (benef_i / Σ abiertas). Credited to 'el mercado' / 'el sistema' they are false (Nueva Masvida is 12,4 % of the
# open Isapres but 299.429 / 2.486.683 = 12,0 % of the system; Esencial 2,5 % vs 2,4 %).
@pytest.mark.parametrize("text,error", [
    ("Para la competencia, Nueva Masvida ya tiene 12,4% del mercado Isapre en agosto.", "share_base_mismatch"),
    ("Para la competencia, Esencial ya tiene 2,5% de participación de mercado en agosto.", "share_base_mismatch"),
    ("Para la competencia, Nueva Masvida ya tiene 12,4% del sistema en agosto.", "share_base_mismatch"),
    ("Para la competencia, Esencial ya tiene 2,5% en agosto.", "share_base_unstated"),
    ("Para la competencia, Esencial ya tiene 2,5% de participación en agosto.", "share_base_unstated"),
    ("Para la competencia, Esencial sumó 0,3 puntos de participación de mercado desde enero.", "share_base_mismatch"),
    ("Para la competencia, Banmédica explica el 77% del mercado desde enero.", "share_base_mismatch"),  # contribution
    # 'puntos entre las abiertas' belongs to the Isapre named in the clause; with none named it is the open group's
    ("Para la competencia, las abiertas sumaron 0,3 puntos de participación desde enero.", "number_entity_mismatch"),
    # 'doce meses' / '12 meses' is a period window, never a figure: no 12-month entry exists in this package
    ("Para la competencia, Nueva Masvida crece 1,9% en doce meses.", "number_period_mismatch"),
    ("Para la competencia, Nueva Masvida crece 1,9% en los últimos 12 meses.", "number_period_mismatch"),
    # figure-free superlatives are rank claims (rank.alza: Esencial 1, Nueva Masvida 2; rank.baja: Banmédica 1)
    ("Para la competencia, Nueva Masvida es la Isapre que más crece desde enero.", "superlative_entity_mismatch"),
    ("Para la competencia, Banmédica es la Isapre que más crece desde enero.", "superlative_entity_mismatch"),
    ("Para la competencia, Nueva Masvida lidera la caída desde enero.", "superlative_entity_mismatch"),
    ("Para la competencia, Banmédica es la Isapre más grande.", "superlative_entity_mismatch"),  # Colmena is
])
def test_fourth_review_probes_fail(root, text, error):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    r = sf.verify_lectura(text, f1.build_package(root), ficha=fi)
    assert r["verdict"] == "FAIL" and any(e.startswith(error) for e in r["errors"]), r
    assert not any(e.startswith("number_") and e.split(":", 1)[1] in ("12", "doce") for e in r["errors"]), r


@pytest.mark.parametrize("text", [
    "Para la competencia, Nueva Masvida ya tiene 12,4% de los beneficiarios entre las Isapres abiertas en agosto.",
    "Para la competencia, Esencial ya tiene 2,5% de participación entre las abiertas en agosto.",
    "Para la competencia, Esencial ya tiene 2,5% de participación entre las Isapres abiertas en agosto.",
    "Para el mercado, Esencial ya tiene 2,5% de participación, entre las abiertas.",  # the audience is not a base
    "Para la competencia, la participación de Esencial entre las abiertas subió 0,3 puntos desde enero.",
    "Para la competencia, Esencial sumó 0,3 puntos de participación entre las abiertas desde enero.",
    "Para la competencia, Nueva Masvida crece 1,9% en el año.",  # 'en el año' is YTD wording
    "Para la competencia, Esencial es la Isapre que más crece desde enero.",
    "Para la competencia, Banmédica es la Isapre que más cae desde enero.",
    "Para la competencia, Banmédica lidera la caída desde enero.",
    "Para la competencia, Esencial y Nueva Masvida son las que más crecen desde enero.",
    "Para la competencia, Esencial es la que más crece en el año.",
])
def test_fourth_review_good_sentences_pass(root, text):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    r = sf.verify_lectura(text, f1.build_package(root), ficha=fi)
    assert r["verdict"] == "PASS", r


def test_twelve_month_window_traces_only_a_twelve_month_entry(root):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    p = f1.build_package(root)
    yoy = {"id": "x.81.12m", "value": 1.9, "unit": "%", "entity": 81, "signed": True, "periods": ["2025-08", "2026-08"],
           "requires_period": False, "scope": "insurer:81", "period_kind": "interval", "base": "stock_base"}
    p12 = {**p, "lectura_pool": [*p["lectura_pool"], yoy], "periods": ["2025-08", *p["periods"]]}
    for text in ("Para la competencia, Nueva Masvida crece 1,9% en doce meses.",
                 "Para la competencia, Nueva Masvida crece 1,9% en los últimos 12 meses."):
        assert sf.verify_lectura(text, p12, ficha=fi)["verdict"] == "PASS", text
    # A streak stays a figure: 'doce meses consecutivos' must trace to the streak rank (2 here), and does not.
    r = sf.verify_lectura("Para la competencia, los cotizantes suben por doce meses consecutivos.", p, ficha=fi)
    assert any(e.startswith("number_") for e in r["errors"]), r


def test_share_and_rate_entries_declare_their_denominator(root):
    p = f1.build_package(root)
    by = {e["id"]: e for e in p["lectura_pool"] + p["derived_ratios"]}
    for e in by.values():
        if e["unit"] in ("%", "pp", "share"):
            assert e.get("base"), e
    assert by["share.81.2026-08"]["base"] == by["share.108.delta_pp"]["base"] == "abiertas"
    assert by["contrib.99"]["base"] == by["cargas.contrib"]["base"] == "variacion_sistema"
    assert by["delta.81.ytd_pct"]["base"] == "stock_base"
    line = sf.lv.describe(by["share.81.2026-08"], {81: "Nueva Masvida"})
    assert "participación entre Isapres abiertas" in line["unidad"] and "abiertas" in line["denominador"]
    with pytest.raises(ValueError):
        sf.pool_entry(p["values"], "share.81.2026-08", entity=81)  # a share without a denominator
    with pytest.raises(ValueError):
        sf.pool_entry(p["values"], "share.81.2026-08", entity=81, base="mercado")
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    for prompt in (sf.lectura_prompt(fi, p), sf.review_prompt(fi, p, "x")):
        assert "participación entre Isapres abiertas" in prompt and "denominador" in prompt


def test_good_lecturas_publish_end_to_end(root):
    for text in GOOD_LECTURAS:
        shutil.rmtree(root / "data/stats_fichas", ignore_errors=True)
        r = f1.run(root, today=TODAY, llm=True, model_call=fake_model(text), budget=ALLOW)
        assert r["ficha"]["lectura"]["status"] == "pass", r["ficha"]["lectura"]
        assert sf.public_ficha(r["ficha"])["lectura"]["text"] == text


def test_pool_entries_declare_scope_period_kind_and_trivial(root):
    p = f1.build_package(root)
    by = {e["id"]: e for e in p["lectura_pool"] + p["derived_ratios"] + p["ranks"]}
    assert by["sistema.cargas.2026-08"]["scope"] == "system" and by["sistema.cargas.2026-08"]["period_kind"] == "stock_at"
    assert by["abiertas.beneficiarios.2026-08"]["scope"] == "open"
    assert by["delta.cerradas.ytd"]["scope"] == "closed" and by["delta.cerradas.ytd"]["period_kind"] == "ytd"
    assert by["delta.99.ytd"]["scope"] == "insurer:99"
    assert by["sistema.beneficiarios.mm.2026-08"]["period_kind"] == "month"
    assert by["movilidad.salidas_intervalo"]["period_kind"] == "interval"
    assert by["conteo.abiertas"]["trivial"] is True and by["conteo.abiertas_crecen"]["trivial"] is True
    assert by["rank.baja.67"]["ranking"] == "baja" and by["rank.tamano.99"]["ranking"] == "tamano"
    for e in by.values():
        assert e["scope"] in ("system", "open", "closed") or e["scope"].startswith("insurer:"), e
        assert e["period_kind"] in ("month", "ytd", "interval", "stock_at"), e
    with pytest.raises(ValueError):
        sf.pool_entry(p["values"], "delta.99.ytd", scope="group")


def test_vida_3_is_one_insurer_and_never_a_number(root):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    p = f1.build_package(root)
    for name in ("Vida 3", "Vida Tres", "VidaTres"):
        r = sf.verify_lectura(f"Para la competencia, {name} mantiene su cartera.", p, ficha=fi)
        assert r["errors"] == ["insurer_not_in_facts:vida_tres"], r  # never number_unmapped:3 or a traced 3
    facts = {**fi["facts"], "paragraphs": [*fi["facts"]["paragraphs"], "Vida Tres registró −991 beneficiarios."]}
    assert sf.verify_lectura("Para la competencia, Vida 3 mantiene su cartera.", p, ficha={**fi, "facts": facts})["verdict"] == "PASS"


def test_insurer_fuzzy_matching_spares_common_words(root):
    cat = sf.insurer_catalog()
    assert sf.insurers_mentioned("Escencial y Banmédika", cat) == {"esencial", "banmedica"}
    assert sf.insurers_mentioned("BMI y CB", cat) == {"banmedica", "cruz_blanca"}
    assert sf.insurers_mentioned("bmi y cb", cat) == set()  # short acronyms only in upper case
    assert sf.insurers_mentioned("Un caso especial de salud en la columna del seguro de vida.", cat) == set()


def test_chart_headline_counts_match_the_data(root):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    p = f1.build_package(root)
    up = next(r["value"] for r in p["derived_ratios"] if r["id"] == "conteo.abiertas_crecen")
    assert up == sum(1 for c in sf.load_registry()["families"]["F1"]["open_codes"] if p["kpis"]["deltas"][c] > 0) == 3
    chart = next(c for c in sf.all_charts(fi) if c["id"] == "f1-isapres")
    assert chart["headline"] == "Una Isapre concentra la baja; tres crecen"  # Consalud +397 grows too
    shown_up = sum(1 for r in chart["rows"] if r["value"] > 0)
    assert shown_up == up
    wrong = json.loads(json.dumps(fi))
    next(c for t in wrong["tabs"] for c in t.get("charts") or [] if c["id"] == "f1-isapres")["headline"] = "Una Isapre concentra la baja; dos crecen"
    assert "data:chart_headline_count_not_in_evidence:f1-isapres:dos" in sf.verify_ficha(wrong, p)["errors"]
    assert sf.count_phrase(1, "crece", "crecen") == "una crece" and sf.count_phrase(4, "crece", "crecen") == "cuatro crecen"


def test_audience_terms_are_one_list(root):
    p = f1.build_package(root)
    r = sf.verify_lectura("Para el mercado, la señal es de estabilidad en la cartera.", p)
    assert "business_lens_no_audience" not in r["errors"]
    assert "mercado" in sf.audience_terms() and "competencia" in sf.audience_terms()


def _with_ratio(p, rid, value, unit, entity=None):
    return {**p, "derived_ratios": [*p["derived_ratios"], {"id": rid, "value": value, "unit": unit, "entity": entity,
                                                          "signed": False, "periods": [p["period"]], "requires_period": False}]}


@pytest.mark.parametrize("text,ratio", [
    ("Para la competencia, la baja de Banmédica es el doble que la de Esencial.", ("r.doble", 2.04, "x", 99)),
    ("Para la competencia, la baja de Banmédica es quíntuple que la de Esencial.", ("r.q", 4.9, "x", 99)),
    ("Para la competencia, uno de cada cinco beneficiarios perdidos por la cartera es de Banmédica.", ("r.f", 0.2, "share", 99)),
    ("Para la competencia, Banmédica cede la mitad de lo que pierde la cartera de las abiertas.", ("r.m", 0.49, "share", 99)),
    ("Para la competencia, Banmédica cede medio punto de participación entre las abiertas.", ("r.pp", -0.5, "pp", 99)),
])
def test_lectura_number_words_pass_only_with_a_precomputed_ratio(root, text, ratio):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    p = f1.build_package(root)
    assert sf.verify_lectura(text, p, ficha=fi)["verdict"] == "FAIL"
    rid, value, unit, entity = ratio
    r = sf.verify_lectura(text, _with_ratio(p, rid, value, unit, entity), ficha=fi)
    assert r["verdict"] == "PASS", r


def test_number_word_counts_trace_but_are_trivial(root):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    p = _with_ratio(f1.build_package(root), "r.c", 2, "count")
    r = sf.verify_lectura("Para inversionistas, un par de Isapres abiertas sostiene la cartera.", p, ficha=fi)
    assert not any(e.startswith("number_") for e in r["errors"]) and any(e.startswith("restates_shown") for e in r["errors"]), r


def test_lectura_ordinals_and_roman_numerals_trace_to_ranks_and_periods(root):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    p = f1.build_package(root)
    ok = sf.verify_lectura("Para la competencia, Nueva Masvida es la segunda Isapre que más crece desde enero.", p, ficha=fi)
    assert ok["verdict"] == "PASS", ok  # rank.alza.81 = 2
    bad = sf.verify_lectura("Para la competencia, Esencial es la segunda Isapre que más crece desde enero.", p, ficha=fi)
    assert bad["verdict"] == "FAIL" and all(e.startswith("number_") for e in bad["errors"]), bad  # Esencial is first, not second
    q3 = {**p, "periods": [*p["periods"], "2026-09"]}  # a package that covers the III quarter
    assert not any(e.startswith("period") for e in sf.verify_lectura("Para las aseguradoras, el III trimestre definirá la cartera.", q3, ficha=fi)["errors"])
    assert any(e.startswith("period_not_in_package") for e in sf.verify_lectura("Para las aseguradoras, el cuarto trimestre definirá la cartera.", p, ficha=fi)["errors"])
    assert sf.verify_lectura("Para las aseguradoras, la cartera crece en 2026 solo en Esencial.", p, ficha=fi)["verdict"] == "PASS"


def test_insurer_names_with_number_words_are_masked_before_parsing(root):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    p = f1.build_package(root)
    r = sf.verify_lectura("Para la competencia, Vida Tres mantiene su cartera.", p, ficha=fi)
    assert r["errors"] == ["insurer_not_in_facts:vida_tres"], r  # an insurer error, never a number error
    facts = {**fi["facts"], "paragraphs": [*fi["facts"]["paragraphs"], "Vida Tres registró −991 beneficiarios."]}
    with_vt = {**fi, "facts": facts}
    assert sf.verify_lectura("Para la competencia, Vida Tres mantiene su cartera.", p, ficha=with_vt)["verdict"] == "PASS"
    r = sf.verify_lectura("Para la competencia, Vida Tres cedió 0,7% de su cartera desde enero.", p, ficha=with_vt)
    assert r["verdict"] == "PASS", r  # delta.80.ytd_pct = −0,70 %
    r = sf.verify_lectura("Para la competencia, VidaTres perdió tres beneficiarios.", p, ficha=with_vt)
    assert any(e.startswith("number_") for e in r["errors"]) and not any("vida" in e for e in r["errors"]), r


def test_sentences_without_figures_must_add_an_implication(root):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    p = f1.build_package(root)
    assert any(e.startswith("no_implication") for e in sf.verify_lectura("La serie se mantiene estable en el período.", p, ficha=fi)["errors"])
    # Shown text without figures: the Regional topic's text (the only tab text the C→B layout shows for F1).
    shown = "Alicanto aún no extrae las cifras regionales: la publicación queda enlazada y su ficha suelta se absorbe aquí."
    assert any(e.startswith("restates_shown_text") for e in sf.verify_lectura(shown, p, ficha=fi)["errors"])
    reworded = "Para las Isapres, Alicanto aún no extrae las cifras regionales, porque la publicación solo queda enlazada."
    assert any(e.startswith("restates_shown_text") for e in sf.verify_lectura(reworded, {**p, "causal_note": "n"}, ficha=fi)["errors"])
    two = ("Para la competencia, Banmédica cedió 0,6 pp de participación entre las abiertas desde enero. "
           "Para los prestadores, el riesgo de la cartera se concentra en un solo asegurador.")
    assert sf.verify_lectura(two, p, ficha=fi)["verdict"] == "PASS"
    mixed = two + " La cartera perdió 814 beneficiarios."
    assert any(e.startswith("restates_shown") for e in sf.verify_lectura(mixed, p, ficha=fi)["errors"])


def test_lectura_pool_is_declared_by_the_builder_not_the_whole_package(root):
    p = f1.build_package(root)
    ids = {e["id"] for e in p["lectura_pool"]}
    assert "share.99.delta_pp" in ids and "contrib.99" in ids and not any(i.startswith("cartera.declarado") for i in ids)
    assert all(e["id"] in p["values"] for e in p["lectura_pool"])
    # A package value outside the pool does not trace, even when it is in the package.
    jan = p["values"]["sistema.beneficiarios.2026-01"]["value"]
    text = f"Para las Isapres, en enero la cartera tenía {jan:,} beneficiarios.".replace(",", ".")
    assert any(e.startswith("number_not_in_pool") for e in sf.verify_lectura(text, p)["errors"])


def test_registry_declares_the_figure_rule():
    cfg = sf.load_registry()["lectura"]
    assert "forbid_numbers" not in cfg and "number_words" not in cfg and cfg["jaccard_max"] == 0.5
    assert "Dirección 28-oct 15:25" in cfg["figure_rule"] and cfg["max_words"] == 60


@pytest.mark.parametrize("text", [
    "La cartera se mueve impulsada por las cargas.", "Las cargas impulsan la baja de la cartera.",
    "La baja se origina en las cargas de la cartera.", "La baja de la cartera tiene su origen en las cargas.",
    "La baja responde al alza de precios en las Isapres.", "La baja de la cartera responde a los precios.",
    "Los precios provocaron la baja de la cartera.", "Los precios generaron salidas de la cartera Isapre.",
    "La baja de la cartera se debió a los precios.", "La cartera cae; esto se explica por los precios.",
    "Los precios explicarían la baja de la cartera.", "Los precios causaron la baja de la cartera.",
    "Las Isapres atribuyen la baja de la cartera a los precios.", "La baja se traduce en menos ingreso para las Isapres.",
    "La baja de la cartera lleva a menos ingreso.", "La baja de la cartera llevó a las Isapres a ajustar.",
    "La baja de la cartera produce menos ingreso.", "La cartera cae porque suben los precios.",
    "La cartera cae, ya que suben los precios.", "La cartera cae por efecto de los precios.",
    # second review (28-oct): preterite, future and infinitive forms, and the synonyms
    "La baja de la cartera se tradujo en menos ingreso.", "La baja de la cartera tradujo en menos ingreso a las Isapres.",
    "La baja de la cartera se traducirá en menos ingreso.", "La baja de la cartera traducirá en menos ingreso.",
    "La baja de la cartera puede deberse a los precios.", "La cartera cae debido a los precios.",
    "La baja de la cartera conlleva menos ingreso.", "La baja conllevó ajustes en la cartera.",
    "Los precios inciden en la cartera.", "Los precios incidieron en la cartera.",
    "Los precios influyen en la cartera.", "Los precios influyeron en la cartera.",
    "La baja repercute en la cartera.", "La baja repercutió en la cartera.",
    "Las cargas contribuyen a la baja de la cartera.", "Las cargas contribuyeron a la baja de la cartera.",
    "Los precios gatillan salidas de la cartera.", "Los precios gatillaron salidas de la cartera.",
    "La ley desencadenó salidas de la cartera.", "La ley desencadena salidas de la cartera.",
    "La ley propició salidas de la cartera.", "La ley propicia salidas de la cartera.",
    "La ley condujo a salidas de la cartera.", "La ley conduce a salidas de la cartera.", "La ley conducirá a salidas.",
    "La ley dio lugar a salidas de la cartera.", "La ley da lugar a salidas de la cartera.",
    "La ley dará lugar a salidas de la cartera.", "La ley ha dado lugar a salidas.", "La ley puede dar lugar a salidas.",
    # third review (28-oct): consecutive and origin connectors
    "La cartera cae, por lo que las Isapres ajustan precios.", "La cartera cae, así que las Isapres ajustan.",
    "La cartera cae, de modo que las Isapres ajustan.", "La cartera cae, de manera que las Isapres ajustan.",
    "La cartera cae, lo que hizo que las Isapres ajusten.", "La cartera cae, lo que llevó a las Isapres a ajustar.",
    "La baja de la cartera nace de los precios.", "La baja de la cartera viene de los precios.",
    "La baja de la cartera viene del alza de precios.", "La baja de la cartera proviene de los precios.",
    "La cartera cae en razón de los precios.", "La cartera cae; por ende, las Isapres ajustan.",
    "La cartera cae; en consecuencia, las Isapres ajustan.",
])
def test_causal_conjugations_rejected_without_note(root, text):
    p = f1.build_package(root)
    errors = sf.verify_lectura(text, p)["errors"]
    assert any(e.startswith("causal_without_note") for e in errors), errors
    assert not any(e.startswith("causal") for e in sf.verify_lectura(text, {**p, "causal_note": "nota"})["errors"])


@pytest.mark.parametrize("text", ["Las Isapres crecen en general.", "La cartera muestra una generación de ingreso estable.",
                                  "La competencia entre Isapres se mantiene.", "Para la cartera, la duda es la rentabilidad.",
                                  "La incidencia de cargas en la cartera es estable.", "La contribución de Banmédica al cambio neto domina.",
                                  "La conducta de la cartera es estable.", "El lugar de Banmédica en la cartera se mantiene.",
                                  # non-causal uses of the third review's connectors
                                  "Por lo general la cartera es estable.", "Así queda la cartera de las Isapres.",
                                  "La razón de cargas por cotizante baja en la cartera.",
                                  "De la misma manera que en julio, la cartera es estable.",
                                  "A su modo, la cartera se mantiene.", "Conviene mirar la cartera de las Isapres.",
                                  "La cartera nacional de las Isapres es estable.", "Por endeudamiento no hay dato en la cartera."])
def test_causal_patterns_do_not_flag_plain_words(root, text):
    assert not any(e.startswith("causal") for e in sf.verify_lectura(text, f1.build_package(root))["errors"])


@pytest.mark.parametrize("text,error", [
    ("Para las Isapres, Fonasa gana espacio frente a la cartera privada.", "insurer_not_in_package:fonasa"),
    ("Para la competencia, Golden Cross mantiene su cartera.", "insurer_not_in_facts:colmena"),
    ("Para la competencia, Colmena mantiene su cartera.", "insurer_not_in_facts:colmena"),
    ("Para la competencia, Consalud mantiene su cartera.", "insurer_not_in_facts:consalud"),
    ("Para la competencia, CruzBlanca mantiene su cartera.", "insurer_not_in_facts:cruz_blanca"),
    ("Para la competencia, Vida Tres mantiene su cartera.", "insurer_not_in_facts:vida_tres"),
    ("Para la competencia, Isalud mantiene su cartera.", "insurer_not_in_facts:isalud"),
    ("Para la competencia, Isapre Fundación mantiene su cartera.", "insurer_not_in_facts:fundacion"),
    ("Para la competencia, Cruz del Norte mantiene su cartera.", "insurer_not_in_facts:cruz_del_norte"),
])
def test_lectura_rejects_insurers_outside_package_or_facts(root, text, error):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    p = f1.build_package(root)
    r = sf.verify_lectura(text, p, ficha=fi)
    assert r["verdict"] == "FAIL" and error in r["errors"], r["errors"]


def test_lectura_accepts_insurers_named_in_facts_with_aliases(root):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    p = f1.build_package(root)
    for text in ("Para la competencia, Banmedica y Nueva Más Vida siguen caminos opuestos en su cartera.",
                 "Para la competencia, Masvida y Esencial ganan cartera frente a Banmédica."):
        assert sf.verify_lectura(text, p, ficha=fi)["verdict"] == "PASS", sf.verify_lectura(text, p, ficha=fi)
    # Without the ficha nothing is in the facts, so no insurer may be named.
    assert "insurer_not_in_facts:banmedica" in sf.verify_lectura("Banmédica pierde cartera.", p)["errors"]
    # Banmédica in the package but removed from the facts → rejected.
    facts = {**fi["facts"], "paragraphs": [x for x in fi["facts"]["paragraphs"] if "Banmédica" not in x]}
    assert "insurer_not_in_facts:banmedica" in sf.verify_lectura("Banmédica pierde cartera.", p, ficha={**fi, "facts": facts})["errors"]


def test_reused_lectura_is_rechecked_with_the_current_verifier(root):
    first = f1.run(root, today=TODAY, llm=True, model_call=fake_model(GOOD), budget=ALLOW)
    first["ficha"]["lectura"]["text"] = "La cartera de las Isapres perdió 814 beneficiarios."  # passed an older verifier
    first["ficha"]["builder_version"] = "old"
    sf.publish(first, root)
    model = fake_model(GOOD)
    r = f1.run(root, today=TODAY, llm=True, model_call=model, budget=ALLOW)
    assert r["ficha"]["lectura"]["status"] == "verifier_fail" and model.calls == []  # no repeat payment
    assert any(e.startswith("restates_shown") for e in r["ficha"]["lectura"]["verification"]["errors"])


def test_lectura_repeating_a_ficha_sentence_is_dropped(root):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    sentence = next(x for tab in fi["tabs"] for x in tab.get("text") or [] if "regionales" in x and not any(ch.isdigit() for ch in x))
    text = "Para las Isapres, la señal es de estabilidad en la cartera. " + sentence
    v = sf.verify_lectura(text, f1.build_package(root), ficha=fi)
    assert any(e.startswith("restates_shown_text") for e in v["errors"]), v  # caught by the Jaccard rule…
    r = f1.run(root, today=TODAY, llm=True, model_call=fake_model(text), budget=ALLOW)
    assert r["status"] == "pass" and r["ficha"]["lectura"]["status"] == "verifier_fail"  # …and the ficha still publishes
    assert sf.public_ficha(r["ficha"])["lectura"] is None
    # The post-attach whole-ficha check (repeated_sentence) still runs as a second gate.
    attached = {**fi, "lectura": {"status": "pass", "text": text}}
    assert any("repeated_sentence" in e for e in sf.verify_ficha(attached, f1.build_package(root))["errors"])


def test_lectura_pass_is_published_and_not_paid_twice(root):
    model = fake_model(GOOD)
    r = f1.run(root, today=TODAY, llm=True, model_call=model, budget=ALLOW)
    assert r["ficha"]["lectura"]["status"] == "pass" and r["ficha"]["lectura"]["text"] == GOOD
    assert model.calls[0]["feature"] == "stats_ficha_lectura" and model.calls[0]["allow_luna_fallback"] is False
    assert sf.verify_ficha(r["ficha"], f1.build_package(root))["verdict"] == "PASS"
    sf.publish(r, root)
    again = f1.run(root, today=TODAY, llm=True, model_call=model, budget=ALLOW)
    assert again["status"] == "no_change" and len(model.calls) == 2  # one writer + one reviewer call, never repeated
    assert r["ficha"]["lectura"]["review"]["verdict"] == "PASS"


def test_lectura_omitted_on_verifier_fail_or_no_budget_ficha_still_publishes(root):
    bad = fake_model("Las Isapres perdieron 123.456 afiliados debido a la ley corta.")
    r = f1.run(root, today=TODAY, llm=True, model_call=bad, budget=ALLOW)
    assert r["status"] == "pass" and r["ficha"]["lectura"]["status"] == "verifier_fail"
    assert sf.public_ficha(r["ficha"])["lectura"] is None
    sf.publish(r, root)
    retry = fake_model(GOOD)
    assert f1.run(root, today=TODAY, llm=True, model_call=retry, budget=ALLOW)["status"] == "no_change"
    assert retry.calls == []  # a paid FAIL is not retried for the same release
    shutil.rmtree(root / "data/stats_fichas")
    blocked = fake_model(GOOD)
    nb = f1.run(root, today=TODAY, llm=True, model_call=blocked, budget=lambda worst: (False, "paced_budget_would_be_exceeded"))
    assert nb["status"] == "pass" and nb["ficha"]["lectura"]["status"] == "no_budget" and blocked.calls == []
    sf.publish(nb, root)
    filled = f1.run(root, today=TODAY, llm=True, model_call=blocked, budget=ALLOW)
    assert filled["status"] == "pass" and filled["ficha"]["lectura"]["status"] == "pass"
    assert len(blocked.writes()) == 1 and len(blocked.reviews()) == 1


# Second-model backstop (third review, item 10): the reviewer is mocked; it never reaches a real model.


def test_review_fail_omits_lectura_and_ficha_still_publishes(root):
    model = fake_model(GOOD, review="FAIL", claim="Esencial y Nueva Masvida consoliden lo ganado")
    r = f1.run(root, today=TODAY, llm=True, model_call=model, budget=ALLOW)
    lect = r["ficha"]["lectura"]
    assert r["status"] == "pass" and lect["status"] == "review_fail" and lect["text"] is None and lect["rejected_text"] == GOOD
    assert lect["review"]["claim"] == "Esencial y Nueva Masvida consoliden lo ganado"
    assert len(model.writes()) == 1 and len(model.reviews()) == 1  # capped at one reviewer call per ficha
    review = model.reviews()[0]
    assert review["allow_luna_fallback"] is False and review["max_output_tokens"] == 200
    assert "Evidencia" in review["prompt"] and GOOD in review["prompt"] and "share.99.delta_pp" in review["prompt"]
    assert sf.public_ficha(r["ficha"])["lectura"] is None and "segundo modelo" in r["ficha"]["machine_note"]
    sf.publish(r, root)
    again = fake_model(GOOD)
    assert f1.run(root, today=TODAY, llm=True, model_call=again, budget=ALLOW)["status"] == "no_change" and again.calls == []


def test_review_without_budget_or_with_error_omits_lectura(root):
    seen = []

    def writer_only(worst):  # the writer gets budget, the reviewer does not
        seen.append(worst)
        return (len(seen) == 1, "ok" if len(seen) == 1 else "budget_deferred")
    model = fake_model(GOOD)
    r = f1.run(root, today=TODAY, llm=True, model_call=model, budget=writer_only)
    assert r["status"] == "pass" and r["ficha"]["lectura"]["status"] == "review_no_budget" and model.reviews() == []
    assert r["ficha"]["lectura"]["review"]["reason"] == "budget_deferred"

    def broken(prompt, **kw):
        if kw["feature"].endswith("_review"):
            raise RuntimeError("weekly_api_budget_exhausted")
        return SimpleNamespace(output={"lectura": GOOD}, model="gpt-5.6-luna", estimated_cost_usd=0.0004, response_id="r")
    shutil.rmtree(root / "data/stats_fichas", ignore_errors=True)
    r = f1.run(root, today=TODAY, llm=True, model_call=broken, budget=ALLOW)
    assert r["status"] == "pass" and r["ficha"]["lectura"]["status"] == "review_error"


def test_reused_lectura_without_review_is_not_published_without_a_model(root):
    first = f1.run(root, today=TODAY, llm=True, model_call=fake_model(GOOD), budget=ALLOW)
    first["ficha"]["lectura"].pop("review")  # a Lectura from before the reviewer existed
    first["ficha"]["builder_version"] = "old"
    sf.publish(first, root)
    r = f1.run(root, today=TODAY, llm=False)
    assert r["ficha"]["lectura"]["status"] == "review_not_run" and sf.public_ficha(r["ficha"])["lectura"] is None


def test_run_cap_covers_writer_and_reviewer(root):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    p = f1.build_package(root)
    reg = sf.load_registry()
    reg["lectura"]["max_run_usd"] = sf.review_worst_usd(fi, p, reg) * 1.01  # room for the reviewer, not for both
    model = fake_model(GOOD)
    out = sf.generate_lectura(fi, p, registry=reg, model_call=model, budget=ALLOW)
    assert out["status"] == "no_budget" and out["reason"] == "run_cap_would_be_exceeded" and model.calls == []


def test_default_budget_respects_paced_monthly_budget(monkeypatch):
    from radar_salud import ai_budget, ai_pricing
    monkeypatch.setattr(ai_budget, "has_capacity", lambda kind: True)
    monkeypatch.setattr(ai_pricing, "month_spend_usd", lambda root=None, now=None: 9.99)
    monkeypatch.setattr(ai_pricing, "paced_budget_usd", lambda now=None: 10.0)
    called = []
    monkeypatch.setattr(ai_budget, "allow_call", lambda kind: called.append(kind) or True)
    assert sf._default_budget(0.05) == (False, "paced_budget_would_be_exceeded") and called == []


def test_verify_ficha_rejects_untraced_numbers(root):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    p = f1.build_package(root)
    tampered = {**fi, "dek": fi["dek"].replace("2.486.683", "2.486.684")}
    assert "number_not_in_evidence:2.486.684" in sf.verify_ficha(tampered, p)["errors"]
    tabs = json.loads(json.dumps(fi["tabs"]))
    tabs[0]["charts"][0]["rows"][0]["value"] += 1
    assert any(e.startswith("chart_row_not_in_evidence") for e in sf.verify_ficha({**fi, "tabs": tabs}, p)["errors"])


def test_four_lenses_reject_redundancy_ui_overload_and_missing_denominator(root):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    p = f1.build_package(root)
    # The facts now live in the tabs: the same sentence in two tabs (or in the dek) is redundant.
    repeated = {**fi, "tabs": [fi["tabs"][0], {**fi["tabs"][1], "text": fi["tabs"][1]["text"] + fi["tabs"][-1]["text"][:1]}, *fi["tabs"][2:]]}
    assert any(e.startswith("editorial:repeated_sentence") for e in sf.verify_ficha(repeated, p)["errors"])
    crowded = {**fi, "tabs": [{**fi["tabs"][0], "charts": fi["tabs"][0]["charts"] * 2}, *fi["tabs"][1:]]}
    assert "ui:tab_over_2_charts:total" in sf.verify_ficha(crowded, p)["errors"]
    no_den = {**fi, "notes": [n for n in fi["notes"] if "denominador" not in n["text"]]}
    assert "data:denominator_not_stated" in sf.verify_ficha(no_den, p)["errors"]
    assert "business:watch_out_of_range" in sf.verify_ficha({**fi, "watch": fi["watch"] * 2}, p)["errors"]
    assert "business_lens_no_audience" in sf.verify_lectura("La serie muestra una trayectoria estable en el período.", p)["errors"]


def test_lectura_prompt_encodes_four_lenses(root):
    fi = f1.run(root, today=TODAY, llm=False)["ficha"]
    prompt = sf.lectura_prompt(fi, f1.build_package(root))
    for lens in ("Negocio de salud", "Analítica de datos", "Editorial", "UI/UX"):
        assert lens in prompt
    assert "Banmédica" in prompt and "Cartera total" in prompt
    # Figure rule: the prompt lists only new material (not what the ficha shows) and asks for an implication.
    assert "Material nuevo" in prompt and "share.99.delta_pp" in prompt and "nunca repitas" in prompt.lower()
    material = sf.lectura_material(fi, f1.build_package(root))
    ids = {x["id"] for x in material["nuevas"] + material["razones"] + material["rankings"]}
    assert "contrib.99" not in ids and "sistema.beneficiarios.mm.2026-08" not in ids and "ratio.salidas_por_entrada" in ids
    assert "Sin cifras" not in prompt and "máximo 60 palabras" in prompt


def test_share_page_has_og_preview_and_deep_link(root, tmp_path):
    sf.publish(f1.run(root, today=TODAY, llm=False), root)
    card = sf.apply_to_snapshot([], sf.live_fichas(root))[0]
    assert card["share_url"] == "data/fichas/f1_cartera.html"
    written = sf.write_share_pages([card], tmp_path / "web/data")
    html = Path(written[0]).read_text(encoding="utf-8")
    assert '<meta property="og:title" content="La cartera Isapre casi deja de caer' in html
    assert '<meta name="twitter:card" content="summary"/>' in html  # no real preview image
    assert 'og:image" content="https://cftorre1.github.io/Radar-salud/assets/' in html
    assert 'url=../../index.html?abrir=1#sig-stats-ficha:F1:2026-08' in html  # opens the detail


# ---------------------------------------------------------------------------------------------
# Snapshot: one live ficha per family


def test_snapshot_replaces_pulse_and_loose_cartera_cards(root):
    sf.publish(f1.run(root, today=TODAY, llm=False), root)
    fichas = sf.live_fichas(root)
    sis = {"source_name": "Superintendencia de Salud", "signal_types": ["Datos"]}
    pulse = {**sis, "title": "Pulso Isapre · datos a 2026-07", "event_type": "DATA_PULSE"}
    loose = {**sis, "title": "Estadística Mensual de Cartera de Beneficiarios del Sistema ISAPRE a Nivel Regional – Agosto 2026", "event_date": "2026-10-02"}
    older = {**sis, "title": "Estadística Mensual de Cartera de Beneficiarios del Sistema ISAPRE a Nivel Regional – Julio 2026", "event_date": "2026-09-07"}
    yearly = {**sis, "title": "Estadística Mensual de Cartera de Beneficiarios del Sistema ISAPRE – año 2026", "event_date": "2026-09-07"}
    keep = {**sis, "title": "Estadística Trimestral de Casos GES"}
    out = sf.apply_to_snapshot([pulse, loose, older, yearly, keep], fichas)
    titles = [c["title"] for c in out]
    for gone in (pulse, loose, older, yearly):
        assert gone["title"] not in titles
    assert keep["title"] in titles
    card = out[-1]
    assert card["stats_ficha"]["period"] == "2026-08" and card["data_period"] == "2026-08"
    assert card["editorial_committee"]["verdict"] == "PASS" and card["piece_id"] == "stats-ficha:F1:2026-08"
    assert "evidence" not in card["stats_ficha"] and card["radar_score"] >= 85
    for key in ("title", "signal_types", "scopes", "event_date", "source_name", "source_url", "card_what"):
        assert card[key]


def test_snapshot_keeps_releases_newer_than_the_ficha(root):
    """Same rule as #84 (6099d0d8): a release the ficha does not cover stays visible until the ficha updates."""
    sf.publish(f1.run(root, today=TODAY, llm=False), root)
    fichas = sf.live_fichas(root)
    sis = {"source_name": "Superintendencia de Salud", "signal_types": ["Datos"]}
    september = {**sis, "title": "Estadística Mensual de Cartera de Beneficiarios del Sistema ISAPRE a Nivel Regional – Septiembre 2026", "event_date": "2026-10-30"}
    by_field = {**sis, "title": "Cartera de beneficiarios regional", "data_period": "2026-09", "event_date": "2026-09-01"}
    yearly_new = {**sis, "title": "Estadística Mensual de Cartera de Beneficiarios del Sistema ISAPRE – año 2026", "event_date": "2026-11-02"}
    mobility_new = {**sis, "title": "Estadística Mensual de Movilidad de Cartera de Cotizantes del Sistema ISAPRE a Nivel Regional – Año 2026", "event_date": "2026-11-02"}
    titles = [c["title"] for c in sf.apply_to_snapshot([september, by_field, yearly_new, mobility_new], fichas)]
    for kept in (september, by_field, yearly_new, mobility_new):
        assert kept["title"] in titles
    assert sf.loose_card_period(september) == "2026-09" and sf.loose_card_period(yearly_new) is None


def test_repo_snapshot_has_august_ficha_and_no_wrong_ges_finding():
    rows = json.loads((REPO / "web/data/radar_today.json").read_text(encoding="utf-8"))["signals"]
    fichas = [s for s in rows if s.get("stats_ficha")]
    assert len(fichas) == 1 and fichas[0]["data_period"] == "2026-08"
    assert not any(str(s.get("title", "")).startswith("Pulso Isapre") for s in rows)
    ges = next(s for s in rows if "casos-ges-auge-de-fonasa-y-sistema-isapre-a-junio-2026" in str(s.get("source_url")))
    assert ges["data_insights"] == [] and ges["statistical_view"]["publication_status"] == "HOLD"
    assert ges["statistical_view"]["correction"]["text"].startswith("Corregido el 28 oct:")
    # The withdrawn sentence survives only inside the correction record, never in a displayed field.
    shown = json.dumps([{k: v for k, v in s.items() if k != "data_insight_meta"} for s in rows], ensure_ascii=False)
    assert "presenta el mayor volumen" not in shown


# ---------------------------------------------------------------------------------------------
# GES card bug: pinned current-year sheet, sum vs declared total, HOLD


def ges_book(rows_2026, extra_sheets=("Año 2005",)):
    from openpyxl import Workbook
    book = Workbook()
    first = True
    for name in extra_sheets:
        ws = book.active if first else book.create_sheet(name)
        ws.title = name
        first = False
        ws.append(["N°", "Problema de salud", "Fonasa", "Isapre"])
        ws.append([21, 21, 900000, 900000])
        ws.append([22, 22, 1, 1])
        ws.append([23, 23, 1, 1])
    ws = book.create_sheet("Año 2026") if not first else book.active
    ws.title = "Año 2026"
    ws.append(["N°", "Problema de salud", "Fonasa", "Isapre"])
    for r in rows_2026:
        ws.append(r)
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


TITLE = "Estadística Trimestral de Casos GES (AUGE) de Fonasa y Sistema ISAPRE – a junio 2026"
ROWS = [[1, "Diabetes", 120, 30], [2, "Cáncer de mama", 80, 40], [3, "Hipertensión", 150, 20]]


def test_ges_pins_current_year_sheet_and_reconciles():
    r = data_insights.parse_ges_workbook(ges_book(ROWS + [[None, "Total", 350, 90]]), TITLE)
    assert r["status"] == "validated" and r["sheet"] == "Año 2026" and r["parser"] == sf.GES_PARSER
    assert r["reconciliation"] == {"ok": True, "rows_sum": [350, 90], "declared_total": [350, 90], "total_row": 5}
    assert r["insights"][0].startswith("Hipertensión") and "21" not in r["insights"][0]
    assert sf.ges_finding_valid(r)


@pytest.mark.parametrize("rows,extra,reason", [
    (ROWS + [[None, "Total", 351, 90]], ("Año 2005",), "HOLD_GES_RECONCILIATION"),
    (ROWS, ("Año 2005",), "HOLD_GES_TOTAL_MISSING"),
    ([[1, "21", 120, 30], [2, "22", 80, 40], [3, "23", 150, 20], [None, "Total", 350, 90]], ("Año 2005",), "HOLD_GES_LABEL_NOT_TEXT"),
])
def test_ges_holds_on_failures(rows, extra, reason):
    r = data_insights.parse_ges_workbook(ges_book(rows, extra), TITLE)
    assert r["status"] == "HOLD" and r["hold_reason"] == reason and r["insights"] == []


def test_ges_holds_when_year_sheet_is_not_unique():
    from openpyxl import Workbook
    book = Workbook()
    book.active.title = "Año 2005"
    stream = BytesIO()
    book.save(stream)
    r = data_insights.parse_ges_workbook(stream.getvalue(), TITLE)
    assert r["status"] == "HOLD" and r["hold_reason"] == "HOLD_GES_SHEET_NOT_PINNED"


WITHDRAWN = "21 presenta el mayor volumen combinado entre las filas comparables detectadas en la tabla GES del corte 2026-06."


def test_old_ges_finding_goes_to_hold_and_keeps_withdrawn_text_in_history():
    old = {"title": "GES", "source_url": "https://x/ges", "data_insights": [WITHDRAWN],
           "data_insight_meta": {"family": "ges", "status": "validated", "sheet": "Año 2005", "insights": [WITHDRAWN]}}
    out = sf.sanitize_ges_signal(old)
    assert out["data_insights"] == [] and out["data_insight_meta"]["status"] == "HOLD"
    sup = out["data_insight_meta"]["superseded_meta"]
    assert sup["withdrawn_insight_texts"] == [WITHDRAWN] and sup["withdrawn_data_insights"] == [WITHDRAWN]
    assert sup["corrected_at"] == "2026-10-28" and sup["reason"] == "parser leía hoja Año 2005"
    assert sup["superseded_parser"] == sf.GES_LEGACY_PARSER and sup["corrected_by_parser"] == sf.GES_PARSER
    view = out["statistical_view"]
    assert view["publication_status"] == "HOLD" and view["correction"]["text"].startswith("Corregido el 28 oct:")
    assert "21 presenta" not in json.dumps(view, ensure_ascii=False)  # shown text never repeats the withdrawn finding
    assert sf.sanitize_ges_signal(out) is out  # idempotent: the record is not rewritten again
    assert view["correction"]["label"] == "28 oct"  # the row tag and the detail note share this label
    old_label = {**out, "statistical_view": {**view, "correction": {"corrected_at": "2026-10-28", "text": "Corregido el 28-oct: x"}}}
    relabeled = sf.sanitize_ges_signal(old_label)["statistical_view"]["correction"]
    assert relabeled["label"] == "28 oct" and relabeled["text"].startswith("Corregido el 28 oct:")
    # An already-held record without the text (the first version of this fix) is upgraded, never emptied.
    legacy = {**old, "data_insights": [], "data_insight_meta": {"family": "ges", "status": "HOLD", "parser": sf.GES_PARSER,
              "superseded_meta": {"family": "ges", "sheet": "Año 2005", "withdrawn_insights": 1}}}
    up = sf.sanitize_ges_signal(legacy)
    assert up["data_insight_meta"]["superseded_meta"]["withdrawn_insights"] == 1
    assert up["statistical_view"]["correction"]["corrected_at"] == "2026-10-28"
    good = {"data_insights": ["x"], "data_insight_meta": {"family": "ges", "status": "validated", "parser": sf.GES_PARSER, "reconciliation": {"ok": True}}}
    assert sf.sanitize_ges_signal(good) is good
    fresh_hold = {"data_insights": [], "data_insight_meta": {"family": "ges", "status": "HOLD", "parser": sf.GES_PARSER}}
    assert sf.sanitize_ges_signal(fresh_hold) is fresh_hold


def test_history_keeps_the_withdrawn_ges_sentence_verbatim():
    rows = json.loads((REPO / "data/history/superintendencia_signals.json").read_text(encoding="utf-8"))["signals"]
    ges = next(s for s in rows if "casos-ges-auge-de-fonasa-y-sistema-isapre-a-junio-2026" in str(s.get("source_url")))
    sup = ges["data_insight_meta"]["superseded_meta"]
    assert sup["withdrawn_insight_texts"][0].startswith("21 presenta el mayor volumen combinado")
    assert sup["reason"] == "parser leía hoja Año 2005" and sup["corrected_at"] == "2026-10-28"


def test_registry_declares_f1_rules():
    reg = sf.load_registry()
    f1cfg = reg["families"]["F1"]
    assert f1cfg["status"] == "active" and len(f1cfg["open_codes"]) == 7
    assert "abiertas" in f1cfg["kpis"]["participacion"]["denominator"].lower()
    assert reg["families"]["F2"]["status"] == "planned" and reg["lectura"]["max_words"] == 60 <= 90  # short Lectura
