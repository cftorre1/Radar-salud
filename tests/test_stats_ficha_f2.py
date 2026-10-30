"""F2 · FEFI (financial statements of the Isapres): structure template, capture, Esencial HOLD and the ficha.

No network and no model calls. Workbooks are synthetic (built here with openpyxl, clearly fixture values) in a layout
like the SIS one: one sheet per statement, Isapres in rows, a two-row header (period over item), a 'Total sistema'
row and an 'en miles de pesos' unit line. Every figure asserted below is recomputed here from the fixture inputs.
"""
import copy
import hashlib
import json
import shutil
import sys
import types
from datetime import date
from io import BytesIO
from pathlib import Path

import pytest

from radar_salud import figure_templates as ft
from radar_salud import stats_capture_fefi as fc
from radar_salud import stats_ficha_f2 as f2
from radar_salud import stats_fichas as sf

REPO = Path(__file__).resolve().parents[1]
TODAY = date(2026, 10, 29)
URL = "https://www.superdesalud.gob.cl/app/uploads/fixture/finan_ifrs_mar_2026_fixture.xlsx"
# code → name in the fixture sheet (the SIS prints names; codes come from data/entities/isapres.json)
ISAPRES = {"67": "Colmena Golden Cross", "78": "Cruz Blanca", "80": "Vida Tres", "81": "Nueva Masvida", "99": "Banmédica",
           "107": "Consalud", "108": "Esencial", "63": "Isalud", "76": "Fundación", "94": "Cruz del Norte"}
M = 1_000_000  # fixture values are in miles de pesos: M miles = 1.000 MM$
DATA = {  # (t = marzo 2026, t_1 = marzo 2025); costs printed as negatives, as many IFRS sheets do
    "resultado": {"t": [12 * M, 8 * M, 1 * M, 2_500_000, 9 * M, 1_500_000, -2 * M, 500_000, 1_090_653, 500_000],
                  "t_1": [-3 * M, 1 * M, 200_000, -500_000, 2 * M, -1 * M, -500_000, 300_000, 1_026_715, 200_000]},
    "ingresos": {"t": [x * M for x in (150, 120, 20, 60, 140, 110, 30, 15, 12, 5)],
                 "t_1": [x * M for x in (140, 118, 19, 58, 135, 108, 25, 14, 11, 5)]},
    "costo": {"t": [-x * M for x in (130, 105, 17, 52, 120, 98, 28, 13, 10, 4)],
              "t_1": [-x * M for x in (125, 104, 17, 51, 118, 97, 24, 12, 10, 4)]},
    "cotizaciones": {"t": [x * M for x in (140, 110, 18, 55, 130, 100, 28, 14, 11, 5)],
                     "t_1": [x * M for x in (130, 108, 17, 53, 125, 99, 23, 13, 10, 5)]},
    "prestaciones": {"t": [-x * M for x in (110, 90, 15, 45, 105, 85, 25, 11, 9, 4)],
                     "t_1": [-x * M for x in (100, 88, 14, 44, 100, 84, 21, 10, 8, 4)]},
}
PATRIMONIO = [x * M for x in (80, 60, 10, 20, 90, 30, -5, 8, 6, 3)]
NOTE = ("Nota 3: la información de Isapre Esencial S.A. corresponde a su primer período de operación y no es comparable "
        "con los registros institucionales (texto de fixture).")
ITEMS = (("ingresos", "Ingresos operacionales"), ("costo", "Costo operacional"), ("cotizaciones", "Ingresos por cotizaciones"),
         ("prestaciones", "Costo de prestaciones de salud"), ("resultado", "Resultado del ejercicio"))


def workbook(data=None, patrimonio=None, *, extra_row=None, total_delta=0, dup_sheet=False, drop=None, text_cell=None,
             no_unit=False, note=NOTE) -> bytes:
    from openpyxl import Workbook
    data = data or DATA
    patrimonio = patrimonio or PATRIMONIO
    wb = Workbook()
    ws = wb.active
    ws.title = "Notas Explicativas"
    for i in range(1, 18):
        ws.cell(row=i, column=1, value=f"Nota general {i} (fixture)." if i % 3 == 0 else None)
    ws["A18"] = note
    for name in (["E. Resultados comparado por Isa"] + (["E. Resultados comparado por Isa (2)"] if dup_sheet else [])):
        ws = wb.create_sheet(name)
        ws.append(["Estado de resultados comparado por Isapre"])
        ws.append([None] if no_unit else ["(en miles de pesos)"])
        ws.append([])
        ws.append(["Código", "Isapre", "Marzo 2026"] + [None] * 4 + ["Marzo 2025"] + [None] * 4)
        ws.append([None, None] + [lab for _, lab in ITEMS] * 2)
        for k, (code, name_) in enumerate(ISAPRES.items()):
            if code == drop:
                continue
            row = [int(code), name_] + [data[i]["t"][k] for i, _ in ITEMS] + [data[i]["t_1"][k] for i, _ in ITEMS]
            if text_cell and code == text_cell:
                row[2] = "s/i"
            ws.append(row)
        if extra_row:
            ws.append([999, extra_row] + [1] * 10)
        ws.append([None, "Total sistema"] + [sum(data[i]["t"]) + total_delta for i, _ in ITEMS]
                  + [sum(data[i]["t_1"]) for i, _ in ITEMS])
        ws.append([None, "Total Isapres abiertas (otra tabla)"] + [1] * 10)
    ws = wb.create_sheet("Balance por Isapre")
    ws.append(["Balance (en miles de pesos)"])
    ws.append(["Código", "Isapre", "Patrimonio marzo 2026"])
    for k, (code, name_) in enumerate(ISAPRES.items()):
        ws.append([int(code), name_, patrimonio[k]])
    ws.append([None, "Total sistema", sum(patrimonio)])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def col(item, period, label, sign=1):
    yr = "{t_year}" if period == "t" else "{t1_year}"
    return {"pattern": rf"^{{t_month}} {yr} {label}$", "item": item, "period": period, **({"sign": sign} if sign != 1 else {})}


def reviewed_template():
    lab = {k: ft.normalize_text(v) for k, v in ITEMS}
    cols = {}
    for key, item, sign in (("resultado", "resultado_ejercicio", 1), ("ingresos", "ingresos_operacionales", 1),
                            ("costo", "costo_operacional", -1), ("cotizaciones", "ingresos_cotizaciones", 1),
                            ("prestaciones", "costo_prestaciones", -1)):
        for p in ("t", "t_1"):
            cols[f"{item}.{p}"] = col(item, p, lab[key], sign)
    unit = {"pattern": r"en miles de pesos", "to_mm_clp": 0.001, "label": "MM$"}
    return {
        "reviewed_by": "Fixture reviewer", "reviewed_at": "2026-10-30",
        "entities": {"67": r"\bcolmena\b", "78": r"\bcruz blanca\b", "80": r"\bvida tres\b", "81": r"\bnueva masvida\b",
                     "99": r"\bbanmedica\b", "107": r"\bconsalud\b", "108": r"\besencial\b", "63": r"\bisalud\b",
                     "76": r"\bfundacion\b", "94": r"\bcruz del norte\b"},
        "tables": {
            "resultados": {"sheet": r"^e\. resultados comparado por isa$", "orientation": "entities_in_rows",
                           "header_row": r"^codigo \| isapre \| {t_month} {t_year} \| {t_month} {t1_year}$", "header_rows": 2,
                           "label_columns": 2, "total_row": r"^total sistema$", "unit": unit, "sum_tolerance": 0, "columns": cols},
            "balance": {"sheet": r"^balance por isapre$", "orientation": "entities_in_rows",
                        "header_row": r"^codigo \| isapre \| patrimonio {t_month} {t_year}$", "header_rows": 1,
                        "label_columns": 2, "total_row": r"^total sistema$", "unit": unit, "sum_tolerance": 0,
                        "columns": {"patrimonio.t": {"pattern": r"^patrimonio {t_month} {t_year}$", "item": "patrimonio", "period": "t"}}}},
        "coverage_note": {"sheet": r"^notas explicativas$", "pattern": r"\besencial\b.*no es comparable"},
    }


def registry(template=None):
    reg = sf.load_registry()
    reg = copy.deepcopy(reg)
    reg["families"]["F2"]["parser_template"] = reviewed_template() if template is None else template
    return reg


def make_root(tmp_path, *, hold=False, template=None) -> Path:
    root = tmp_path / "root"
    (root / "config").mkdir(parents=True)
    (root / "data/entities").mkdir(parents=True)
    (root / "data/history").mkdir(parents=True)
    shutil.copy(REPO / "data/entities/isapres.json", root / "data/entities/isapres.json")
    rules = {"holds": {"global_ai_coding": "HOLD_SOURCE_METHOD: fixture"}}
    if hold:
        rules["holds"]["fefi_2026_03"] = "HOLD_COVERAGE_ESENCIAL: Notas Explicativas A18 contradicts institutional rows (fixture)."
    (root / "config/editorial_rules_v3.json").write_text(json.dumps(rules), encoding="utf-8")
    (root / "config/stats_fichas_v1.json").write_text(json.dumps(registry(template), ensure_ascii=False), encoding="utf-8")
    signals = {"signals": [{"title": "Estadísticas Financieras del Sistema ISAPRE a marzo 2026", "source_name": "Superintendencia de Salud",
                            "source_url": "https://www.superdesalud.gob.cl/biblioteca-digital/fixture-fefi-marzo-2026/",
                            "event_date": "2026-07-07", "source_documents": [{"url": URL}]}]}
    (root / "data/history/superintendencia_signals.json").write_text(json.dumps(signals), encoding="utf-8")
    return root


def fetcher(data: bytes):
    calls = []

    def fetch(url):
        calls.append(url)
        return data
    fetch.calls = calls
    return fetch


def capture(root, data=None, template=None):
    reg = json.loads((root / "config/stats_fichas_v1.json").read_text(encoding="utf-8"))
    tpl = reg["families"]["F2"]["parser_template"] if template is None else template
    fetch = fetcher(data or workbook())
    fc.capture_fefi(root, fetch=fetch, template=tpl)
    return fetch


def run(root, template=None):
    return f2.run(root, today=TODAY, registry=registry(template), llm=False)


@pytest.fixture(autouse=True)
def isolated_ledgers(tmp_path, monkeypatch):
    monkeypatch.setenv("RADAR_ROOT", str(tmp_path / "ledger-root"))
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)


# ---------------------------------------------------------------------------------------------
# Committed state: the template is a placeholder and the Esencial HOLD is active


def test_committed_registry_keeps_f2_planned_unreviewed_and_on_hold():
    cfg = sf.load_registry()["families"]["F2"]
    assert cfg["status"] == "planned" and cfg["lectura_enabled"] is False
    assert ft.review_state(cfg["parser_template"]) == "placeholder"
    assert set(cfg["parser_template"]["tables"]) >= {"resultados"}
    assert all(c["item"] in fc.ITEMS and c["period"] in fc.PERIODS
               for t in cfg["parser_template"]["tables"].values() for c in t["columns"].values())
    assert f2.esencial_hold(REPO) and f2.esencial_hold(REPO).startswith("fefi_2026_03: HOLD_COVERAGE_ESENCIAL")


def test_committed_placeholder_template_parses_nothing_and_holds(tmp_path):
    committed = sf.load_registry()["families"]["F2"]["parser_template"]
    root = make_root(tmp_path, template=committed)
    capture(root)
    doc = json.loads((root / fc.STORE).read_text(encoding="utf-8"))["releases"][0]["document"]
    assert doc["status"] == "captured" and doc["parse"]["status"] == "template_placeholder"
    assert not doc["parse"].get("tables") and doc["parse"]["sheet_names"][0] == "Notas Explicativas"
    res = run(root, template=committed)
    assert res["status"] == "hold" and any("parser_template_placeholder" in e for e in res["errors"])


@pytest.mark.parametrize("patch,state", [
    ({"reviewed_by": "pendiente"}, "unreviewed"), ({"reviewed_by": ""}, "unreviewed"), ({"reviewed_by": None}, "unreviewed"),
    ({"reviewed_at": None}, "unreviewed"), ({"reviewed_at": ""}, "unreviewed"),
    ({"entities": {**reviewed_template()["entities"], "67": "<<colmena>>"}}, "placeholder")])
def test_unreviewed_or_placeholder_template_holds(tmp_path, patch, state):
    tpl = {**reviewed_template(), **patch}
    root = make_root(tmp_path, template=tpl)
    capture(root)
    doc = json.loads((root / fc.STORE).read_text(encoding="utf-8"))["releases"][0]["document"]
    assert doc["parse"]["status"] == f"template_{state}"
    res = run(root, template=tpl)
    assert res["status"] == "hold" and any(f"parser_template_{state}" in e for e in res["errors"])
    assert not (root / "data/stats_fichas/live/f2_financieras.json").exists()


def test_esencial_hold_disables_aggregates_and_rankings_and_holds(tmp_path):
    root = make_root(tmp_path, hold=True)
    capture(root)
    pkg = f2.build_package(root, registry())
    assert pkg["reconciliation"]["ok"] is False
    assert any(x.startswith("HOLD_COVERAGE_ESENCIAL:fefi_2026_03") for x in pkg["reconciliation"]["failures"])
    assert pkg["disabled_by_esencial_hold"] == ["agregados del sistema", "rankings entre Isapres"]
    assert not any(".sistema." in k for k in pkg["values"]) and pkg["ranks"] == []  # never computed
    res = run(root)
    assert res["status"] == "hold" and any("HOLD_COVERAGE_ESENCIAL" in e for e in res["errors"])


def test_unreadable_editorial_rules_fail_closed(tmp_path):
    root = make_root(tmp_path)
    (root / "config/editorial_rules_v3.json").write_text("{", encoding="utf-8")
    assert "HOLD_COVERAGE_ESENCIAL" in f2.esencial_hold(root)


# ---------------------------------------------------------------------------------------------
# Parser: reviewed template, exactly-once anchors, sums


def test_reviewed_template_reads_every_figure_with_its_cell_and_reconciles(tmp_path):
    root = make_root(tmp_path)
    capture(root)
    doc = json.loads((root / fc.STORE).read_text(encoding="utf-8"))["releases"][0]["document"]
    p = doc["parse"]
    assert p["status"] == "verified", p.get("failures")
    assert p["template_sha256"] == ft.templates_sha256(reviewed_template()) and p["period"] == "2026-03"
    res = p["tables"]["resultados"]
    assert res["sheet"] == "E. Resultados comparado por Isa" and res["header_row"] == 4
    assert res["values"]["resultado_ejercicio.t"]["total"] == pytest.approx(34090.653)
    assert res["values"]["resultado_ejercicio.t_1"]["total"] == pytest.approx(-273.285)
    assert res["values"]["costo_operacional.t"]["67"] == pytest.approx(130_000)  # printed −130.000.000 miles, sign −1
    assert res["cells"]["resultado_ejercicio.t"]["67"] == "G6" and res["cells"]["resultado_ejercicio.t_1"]["total"] == "L16"
    assert res["rows"]["total"] == 16  # the first 'Total sistema' after the header, not the later 'Total Isapres abiertas'
    assert p["coverage_note"]["cell"] == "A18" and "Esencial" in p["coverage_note"]["text"]
    assert p["tables"]["balance"]["values"]["patrimonio.t"]["108"] == pytest.approx(-5000)
    assert len(json.dumps(doc)) < 40_000 and "PK" not in json.dumps(doc)  # values and hashes, never the binary


@pytest.mark.parametrize("kwargs,why", [
    ({"extra_row": "Isapre nueva"}, "resultados:unlisted_row"),
    ({"total_delta": 5}, "resultados:sum_mismatch"),
    ({"drop": "81"}, "resultados:entity_no_match:81"),
    ({"text_cell": "99"}, "resultados:non_numeric"),
    ({"no_unit": True}, "resultados:unit_no_match"),
    ({"note": "Nota sin la aseguradora."}, "coverage_note:no_match"),
])
def test_parser_fails_closed_on_any_structural_doubt(kwargs, why):
    p = fc.parse_fefi_workbook(workbook(**kwargs), reviewed_template(), "2026-03")
    assert p["status"] == "failed" and any(f.startswith(why) for f in p["failures"]), p["failures"]
    assert p["tables"] == {}


def test_a_pattern_matching_two_sheets_holds_and_an_anchored_one_does_not():
    loose = reviewed_template()
    loose["tables"]["resultados"]["sheet"] = r"resultados comparado por isa"
    p = fc.parse_fefi_workbook(workbook(dup_sheet=True), loose, "2026-03")
    assert "resultados:sheet_multiple_matches:2" in p["failures"]
    assert fc.parse_fefi_workbook(workbook(dup_sheet=True), reviewed_template(), "2026-03")["status"] == "verified"


def test_template_tokens_follow_the_release_period_and_never_match_another_year():
    assert fc.parse_fefi_workbook(workbook(), reviewed_template(), "2026-03")["status"] == "verified"
    other = fc.parse_fefi_workbook(workbook(), reviewed_template(), "2026-06")  # 'junio 2026' is not in this workbook
    assert other["status"] == "failed" and any("header_no_match" in f for f in other["failures"])
    assert ft.compile_pattern("{t_year} x")[1] == "invalid"  # a token without a period never compiles


def test_invalid_template_patterns_hold():
    tpl = reviewed_template()
    tpl["tables"]["resultados"]["columns"]["resultado_ejercicio.t"]["pattern"] = "([unclosed"
    tpl["tables"]["balance"]["columns"]["patrimonio.t"]["item"] = "margen_inventado"
    p = fc.parse_fefi_workbook(workbook(), tpl, "2026-03")
    assert "resultados:column_pattern_invalid:resultado_ejercicio.t" in p["failures"]
    assert "balance:column_spec_invalid:patrimonio.t" in p["failures"]


def test_legacy_xls_needs_xlrd_and_reads_through_it(monkeypatch):
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64
    monkeypatch.setitem(sys.modules, "xlrd", None)  # import fails
    assert fc.parse_fefi_workbook(ole, reviewed_template(), "2026-03")["status"] == "xls_reader_unavailable"
    from openpyxl import load_workbook
    real = load_workbook(BytesIO(workbook()), data_only=True)

    class Sheet:
        def __init__(self, ws):
            self.name, self._rows = ws.title, [list(r) for r in ws.iter_rows(values_only=True)]
            self.nrows = len(self._rows)

        def row_values(self, i):
            return ["" if v is None else v for v in self._rows[i]]

    class Book:
        def __init__(self):
            self._s = [Sheet(ws) for ws in real.worksheets]
            self.nsheets = len(self._s)

        def sheet_by_index(self, i):
            return self._s[i]
    fake = types.SimpleNamespace(open_workbook=lambda file_contents, on_demand: Book())
    monkeypatch.setitem(sys.modules, "xlrd", fake)
    p = fc.parse_fefi_workbook(ole, reviewed_template(), "2026-03")
    assert p["format"] == "xls" and p["status"] == "verified", p.get("failures")


# ---------------------------------------------------------------------------------------------
# Capture: once, hash-bound re-check when the template changes


def test_template_change_triggers_one_hash_bound_recheck_and_keeps_the_previous_parse(tmp_path):
    committed = sf.load_registry()["families"]["F2"]["parser_template"]
    root = make_root(tmp_path, template=committed)
    data = workbook()
    first = capture(root, data, template=committed)
    again = capture(root, data, template=committed)
    assert len(first.calls) == 1 and len(again.calls) == 0  # downloaded once
    rechecked = capture(root, data, template=reviewed_template())
    doc = json.loads((root / fc.STORE).read_text(encoding="utf-8"))["releases"][0]["document"]
    assert len(rechecked.calls) == 1 and doc["parse"]["status"] == "verified" and doc["rechecked_at"]
    assert doc["previous_parse"]["status"] == "template_placeholder"
    assert doc["sha256"] == hashlib.sha256(data).hexdigest()
    assert len(capture(root, data, template=reviewed_template()).calls) == 0  # same template: no further download
    # Different bytes on a re-check: HOLD, never re-captured silently.
    tpl2 = {**reviewed_template(), "reviewed_at": "2026-10-31"}
    capture(root, workbook(total_delta=0, note=NOTE + " Revisado."), template=tpl2)
    doc = json.loads((root / fc.STORE).read_text(encoding="utf-8"))["releases"][0]["document"]
    assert doc["status"] == "hash_mismatch" and doc["recheck"]["recorded_sha256"] == hashlib.sha256(data).hexdigest()
    res = run(root, template=tpl2)
    assert res["status"] == "hold" and any("sha256_mismatch_on_recheck" in e for e in res["errors"])


def test_parse_made_with_another_template_holds(tmp_path):
    root = make_root(tmp_path)
    capture(root)
    other = {**reviewed_template(), "reviewed_at": "2026-11-02"}
    res = run(root, template=other)
    assert res["status"] == "hold" and any("parse_made_with_another_template" in e for e in res["errors"])


def test_capture_only_talks_to_the_sis_fetcher_by_default(tmp_path, monkeypatch):
    root = make_root(tmp_path)
    seen = []

    class Fake:
        def __init__(self, **kw):
            seen.append(kw)

        def __call__(self, url):
            raise PermissionError("robots_disallowed: fixture")
    monkeypatch.setattr(fc.cap, "SisFetcher", Fake)
    fc.capture_fefi(root, template=reviewed_template())
    doc = json.loads((root / fc.STORE).read_text(encoding="utf-8"))["releases"][0]["document"]
    assert seen and doc["status"] == "fetch_failed" and "robots_disallowed" in doc["error"]


# ---------------------------------------------------------------------------------------------
# The ficha (with the Esencial HOLD lifted and the template reviewed)


def test_f2_publishes_adaptive_topics_with_loss_to_profit_rule(tmp_path):
    root = make_root(tmp_path)
    capture(root)
    res = run(root)
    assert res["status"] == "pass", res.get("errors")
    fi = res["ficha"]
    assert [t["id"] for t in fi["tabs"]] == ["resultados", "ingresos_costos", "siniestralidad", "patrimonio"]
    assert fi["headline"] == "Isapres a marzo 2026: utilidad de 34.091 MM$, tras una pérdida de 273 MM$ en marzo 2025"
    # The dek adds what the KPI cards do not show: 577 / 662 and 562 / 633; the siniestralidad spread across Isapres.
    assert fi["dek"] == ("El costo operacional absorbió el 87,2% de los ingresos operacionales (88,8% en marzo 2025). "
                         "Por Isapre, la siniestralidad va de 78,6% (Colmena) a 89,3% (Esencial).")
    k = {x["id"]: x for x in fi["kpis"]}
    assert k["resultado"]["value"] == "34.091 MM$" and "N/A" in k["resultado"]["lines"][1]
    assert k["resultado"]["lines"][1].startswith("+34.364 MM$")  # 34.090,653 − (−273,285) = 34.363,938
    ev = fi["evidence"]["values"]
    assert "resultado_ejercicio.sistema.pct" not in ev  # no % from a negative base
    assert ev["resultado_ejercicio.67.delta"]["value"] == pytest.approx(15_000)
    assert "resultado_ejercicio.67.pct" not in ev and ev["resultado_ejercicio.78.pct"]["value"] == pytest.approx(700)
    # Siniestralidad: costo de prestaciones / ingresos por cotizaciones of the same YTD period (denominator explicit).
    assert ev["siniestralidad.sistema.2026-03"]["value"] == pytest.approx(sum(-x for x in DATA["prestaciones"]["t"]) / sum(DATA["cotizaciones"]["t"]) * 100)
    sin = next(t for t in fi["tabs"] if t["id"] == "siniestralidad")
    assert "Denominador" in sin["charts"][0]["source"] and k["siniestralidad"]["lines"][0] == "de los ingresos por cotizaciones"
    res_tab = fi["tabs"][0]
    assert res_tab["index_line"] == "34.091 MM$ a marzo · 9 de 10 Isapres con utilidad"  # all but Esencial
    assert res_tab["charts"][0]["rows"][0]["label"] == "Mar 2025"  # levels first, the change is in the data table
    table = res_tab["charts"][1]["table"]["rows"]
    assert table[-1][0] == "Sistema" and table[-1][-1] == "N/A"
    assert table[0][0] == "Colmena" and table[-2][0] == "Esencial"  # ordered by result; ties by name
    assert [r[0] for r in table[:-1]].index("Cruz del Norte") < [r[0] for r in table[:-1]].index("Isalud")
    assert fi["verification"]["verdict"] == "PASS" and fi["lectura"]["status"] == "not_attempted"
    assert fi["sources"][0]["url"] == URL and sf.HEX64.fullmatch(fi["sources"][0]["sha256"])
    for text in sf.displayed_texts(fi):
        assert not sf.SIGNED_WITH_WORD.search(text), text


def test_topics_follow_what_the_template_maps(tmp_path):
    tpl = reviewed_template()
    del tpl["tables"]["balance"]
    root = make_root(tmp_path, template=tpl)
    capture(root, template=tpl)
    res = run(root, template=tpl)
    assert res["status"] == "pass" and [t["id"] for t in res["ficha"]["tabs"]] == ["resultados", "ingresos_costos", "siniestralidad"]
    for key in [k for k in tpl["tables"]["resultados"]["columns"] if k.startswith(("ingresos_cot", "costo_prest"))]:
        del tpl["tables"]["resultados"]["columns"][key]
    root2 = make_root(tmp_path / "b", template=tpl)
    capture(root2, template=tpl)
    res = run(root2, template=tpl)
    assert res["status"] == "hold" and any("topics_under_3:resultados,ingresos_costos" in e for e in res["errors"])


def test_threshold_is_evaluated_on_the_rounded_value_shown(tmp_path):
    data = copy.deepcopy(DATA)
    data["cotizaciones"]["t"] = [1000] * 10
    data["prestaciones"]["t"] = [-900] * 9 + [-896]  # 8.996 / 10.000 = 89,96% → shown '90,0%'
    data["cotizaciones"]["t_1"] = [1000] * 10
    data["prestaciones"]["t_1"] = [-890] * 10  # 89,0%
    root = make_root(tmp_path)
    capture(root, workbook(data))
    pkg = f2.build_package(root, registry())
    assert any(f.get("kpi") == "siniestralidad" and f.get("level") == 90 for f in pkg["findings"])
    assert pkg["values"]["siniestralidad.sistema.pp"]["value"] == pytest.approx(1.0)  # 90,0 − 89,0 as shown


@pytest.mark.parametrize("prev,cur,expected,pct", [
    (-273.285, 34090.653, "utilidad de 34.091 MM$, tras una pérdida de 273 MM$ en marzo 2025", None),
    (500.0, -200.0, "pérdida de 200 MM$, tras una utilidad de 500 MM$ en marzo 2025", None),
    (-500.0, -200.0, "pérdida de 200 MM$, menor que la de 500 MM$ en marzo 2025", None),
    (-200.0, -500.0, "pérdida de 500 MM$, mayor que la de 200 MM$ en marzo 2025", None),
    (1000.0, 968.0, "utilidad de 968 MM$, 3,2% menos que en marzo 2025", -3.2),
    (1000.0, 1032.0, "utilidad de 1.032 MM$, 3,2% más que en marzo 2025", 3.2),
])
def test_result_wording_follows_the_sign(prev, cur, expected, pct):
    assert f2.result_phrase(prev, cur, "marzo 2025") == expected
    c = f2.change(prev, cur, result=True)
    assert (c["pct"] is None) == (pct is None) and (pct is None or round(c["pct"], 1) == pct)
    assert not sf.SIGNED_WITH_WORD.search(expected)


def test_verify_ficha_rejects_a_signed_figure_with_a_direction_word(tmp_path):
    root = make_root(tmp_path)
    capture(root)
    fi = run(root)["ficha"]
    pkg = f2.build_package(root, registry())
    bad = copy.deepcopy(fi)
    bad["watch"][0]["explanation"] = "La siniestralidad fue −3,2% menos que el año anterior."
    errs = sf.verify_ficha(bad, pkg)["errors"]
    assert any(e.startswith("signed_figure_with_direction_word:−3,2% menos") for e in errs)


def test_f2_never_calls_a_model_while_lectura_is_disabled(tmp_path):
    root = make_root(tmp_path)
    capture(root)
    calls = []
    res = f2.run(root, today=TODAY, registry=registry(), llm=True, model_call=lambda *a, **k: calls.append(1),
                 budget=lambda w: (True, "ok"))
    assert res["status"] == "pass" and calls == [] and res["ficha"]["lectura"]["status"] == "not_attempted"


# ---------------------------------------------------------------------------------------------
# Shared figure-template helpers


def test_figure_template_review_rules():
    assert ft.review_state(None) == "not_configured"
    assert ft.review_state({"pattern": "<<x>>", "reviewed_by": "a", "reviewed_at": "b"}) == "placeholder"
    assert ft.review_state({"pattern": "x", "_doc": "use <<…>> markers", "reviewed_by": "a", "reviewed_at": "b"}) == "ok"
    assert ft.review_state({"pattern": "x", "reviewed_by": "TODO", "reviewed_at": "b"}) == "unreviewed"
    assert ft.normalize_text("Total  General\n1.004.982 Ñuñoa") == "total general 1004982 nunoa"
    pat, _ = ft.compile_pattern(r"\btotal\b")
    assert ft.match_once(pat, [(1, "total"), (2, "x")])[2] == "ok"
    assert ft.match_once(pat, [(1, "total"), (2, "total")])[2] == "multiple_matches:2"
    assert ft.match_once(pat, [(1, "x")])[2] == "no_match"
    assert ft.compile_pattern("(a)(b)", groups=1)[1] == "invalid"


def test_playbook_is_linked_and_its_code_references_exist():
    import importlib
    import re
    playbook = (REPO / "data/stats_fichas/PLAYBOOK.md").read_text(encoding="utf-8")
    assert "PLAYBOOK.md" in (REPO / "data/stats_fichas/POLICY.md").read_text(encoding="utf-8")
    mods = {"sf": "stats_fichas", "stats_capture": "stats_capture", "stats_capture_fefi": "stats_capture_fefi",
            "figure_templates": "figure_templates", "stats_ficha_f2": "stats_ficha_f2", "lectura_verifier": "lectura_verifier"}
    refs = re.findall(r"`(sf|stats_capture_fefi|stats_capture|figure_templates|stats_ficha_f2|lectura_verifier)\.([A-Za-z_]+)", playbook)
    assert len(refs) >= 20
    for mod, name in refs:
        if mod == "stats_capture" and name == "_recheck_doc":
            continue  # F6 re-check lands with #112 (marked as such in the playbook)
        assert hasattr(importlib.import_module(f"radar_salud.{mods[mod]}"), name), f"{mod}.{name}"
    for path in re.findall(r"`((?:tests|src|scripts|config|web|data)/[\w/.\-*]+)`", playbook):
        if "*" in path or path.endswith("test_stats_fichas_f6_real_pdf.py"):
            continue  # globs; the F6 real-PDF test lands with #112
        assert (REPO / path).exists(), path
