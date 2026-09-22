from __future__ import annotations
import io, re, csv
from urllib.request import Request, urlopen
from typing import Any

def _download(url: str, timeout: int = 30) -> bytes:
    req=Request(url,headers={"User-Agent":"RadarSaludBot/0.2 (+public-source-monitor)"})
    with urlopen(req,timeout=timeout) as r:
        return r.read()

def _sentences(text: str):
    clean=" ".join((text or "").split())
    return [s.strip() for s in re.split(r"(?<=[\.\!\?])\s+",clean) if len(s.strip())>25]

def _extract_validity(text: str) -> dict[str, Any]:
    clean=" ".join((text or "").split())
    patterns=[
        r"((?:entrar[aá]\s+en\s+vigencia|entra\s+en\s+vigencia|regir[aá]\s+desde|rige\s+desde|vigencia)[^.]{0,220}\.)",
        r"((?:a\s+contar\s+de|desde)\s+el?\s*\d{1,2}\s+de\s+[A-Za-zÁÉÍÓÚáéíóúñÑ]+\s+de\s+20\d{2}[^.]{0,180}\.)",
        r"((?:a\s+contar\s+de|desde)\s+la\s+fecha\s+de\s+su\s+publicaci[oó]n[^.]{0,180}\.)",
    ]
    for pat in patterns:
        m=re.search(pat,clean,re.I)
        if m:
            return {"validity_text":m.group(1).strip()}
    return {"validity_text":None}

def pdf_intelligence(url: str) -> dict[str, Any]:
    try:
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(_download(url)))
        text=" ".join((p.extract_text() or "") for p in reader.pages[:18])
    except Exception:
        return {}
    sents=_sentences(text)
    kws=("instruye","modifica","deberá","deberán","prohíbe","establece","vigencia","plazo","resuelve","acoge","rechaza","suspende","cobertura","reportar","remitir","informar")
    ranked=[s for s in sents if any(k in s.lower() for k in kws)]
    key_points=[]
    for s in ranked:
        if s not in key_points:
            key_points.append(s[:360])
        if len(key_points)>=3:break
    t=text.lower(); risks=[]
    if any(x in t for x in ("deberá","deberán","obligación","instruye")):
        risks.append("Revisar obligaciones concretas, responsables internos y evidencia de cumplimiento.")
    if any(x in t for x in ("plazo","vigencia","a contar de","entra en vigencia","rige desde")):
        risks.append("Revisar plazos, fecha de entrada en vigencia y eventuales períodos de transición.")
    if any(x in t for x in ("archivo","reportar","remitir","enviar","información","informar")):
        risks.append("Revisar requerimientos de información, reportabilidad y eventuales ajustes de sistemas o procesos.")
    if any(x in t for x in ("suspende","acoge","rechaza","recurso","reposición")):
        risks.append("Verificar si la resolución altera, confirma o suspende efectos de instrucciones previas.")
    if any(x in t for x in ("cobertura","bonificación","prestación","beneficiario")):
        risks.append("Revisar posibles efectos sobre cobertura, prestaciones y experiencia de beneficiarios.")
    out={"key_points":key_points[:3],"risk_notes":risks[:3]}
    out.update(_extract_validity(text))
    return out

def _period_token(v):
    if v is None:return None
    s=str(v).strip().lower()
    months=("ene","feb","mar","abr","may","jun","jul","ago","sep","oct","nov","dic")
    if re.search(r"\b20\d{2}\b",s) and (any(m in s for m in months) or re.fullmatch(r"20\d{2}",s)):
        return s
    if re.fullmatch(r"\d{1,2}[/-]\d{4}",s):return s
    return None

def spreadsheet_insights(url: str) -> list[str]:
    try:
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(_download(url)),read_only=True,data_only=True)
    except Exception:
        return []
    insights=[]
    for ws in wb.worksheets[:4]:
        rows=[]
        for row in ws.iter_rows(min_row=1,max_row=250,values_only=True):
            vals=list(row[:40])
            if any(v not in (None,"") for v in vals): rows.append(vals)
        if len(rows)<4:continue
        # Pattern A: periods in first column, numeric measures in later columns.
        period_rows=[(i,_period_token(r[0] if r else None)) for i,r in enumerate(rows) if r]
        period_rows=[x for x in period_rows if x[1]]
        if len(period_rows)>=2:
            (i1,p1),(i2,p2)=period_rows[-2],period_rows[-1]
            r1,r2=rows[i1],rows[i2]
            best=None
            for j in range(1,min(len(r1),len(r2))):
                a,b=r1[j],r2[j]
                if isinstance(a,(int,float)) and isinstance(b,(int,float)) and a not in (0,None):
                    pct=(b-a)/abs(a)*100
                    if abs(pct)<=1000 and (best is None or abs(pct)>abs(best[0])):
                        best=(pct,j,a,b)
            if best:
                pct,j,a,b=best
                header=""
                for h in rows[:min(15,len(rows))]:
                    if j<len(h) and isinstance(h[j],str) and h[j].strip():
                        header=h[j].strip()
                label=f" en {header}" if header else ""
                insights.append(f"{ws.title}: entre {p1} y {p2}, la mayor variación detectada{label} fue de {pct:+.1f}%.")
        # Pattern B: period columns in a header row; compare last two period columns across named rows.
        for hi,h in enumerate(rows[:20]):
            pcs=[(j,_period_token(v)) for j,v in enumerate(h)]
            pcs=[x for x in pcs if x[1]]
            if len(pcs)>=2:
                (j1,p1),(j2,p2)=pcs[-2],pcs[-1]
                best=None
                for r in rows[hi+1:]:
                    if len(r)<=max(j1,j2):continue
                    a,b=r[j1],r[j2]
                    label=str(r[0]).strip() if r and r[0] not in (None,"") else ""
                    if label and isinstance(a,(int,float)) and isinstance(b,(int,float)) and a not in (0,None):
                        pct=(b-a)/abs(a)*100
                        if abs(pct)<=1000 and (best is None or abs(pct)>abs(best[0])):
                            best=(pct,label)
                if best:
                    insights.append(f"{ws.title}: {best[1]} mostró la mayor variación entre {p1} y {p2}: {best[0]:+.1f}%.")
                break
        if len(insights)>=3:break
    # Deduplicate and stay conservative.
    out=[]
    for x in insights:
        if x not in out:out.append(x)
    return out[:3]

def analyze_attachments(attachments: list[dict]) -> dict[str, Any]:
    out={"key_points":[],"risk_notes":[],"data_insights":[],"validity_text":None}
    for a in attachments[:6]:
        url=(a or {}).get("url","")
        low=url.lower()
        if ".pdf" in low:
            d=pdf_intelligence(url)
            out["key_points"].extend(d.get("key_points",[]))
            out["risk_notes"].extend(d.get("risk_notes",[]))
            if d.get("validity_text") and not out.get("validity_text"):
                out["validity_text"]=d.get("validity_text")
        elif any(x in low for x in (".xlsx",".xls",".csv")):
            out["data_insights"].extend(spreadsheet_insights(url))
    for k in ("key_points","risk_notes","data_insights"):
        ded=[]
        for x in out[k]:
            if x and x not in ded:ded.append(x)
        out[k]=ded[:3]
    return out
