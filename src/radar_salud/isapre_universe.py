"""Official-code identities, explicit attribution and comparable insurer stocks."""
from __future__ import annotations
import io
import json
import re
from datetime import date
from openpyxl import load_workbook
from .official_tables import read, fold
from .paths import project_root
from .models import EntityRef, MetricObservation

REGISTRY_URL = 'https://www.superdesalud.gob.cl/registro/isapres/'
EXPECTED_CODES = {63, 67, 76, 78, 80, 81, 94, 99, 107, 108}
MONTHS = ('enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre')


def parse_registry(html, *, fetched_at, sha256, overlays=None):
    overlays = overlays or json.loads((project_root() / 'config/isapre_identity_overlays.json').read_text())
    entities = []
    for row in read(html).rows:
        if len(row) < 3 or not row[0].isdigit(): continue
        code = int(row[0])
        if code not in EXPECTED_CODES: raise ValueError(f'Unexpected official Isapre code {code}')
        overlay = overlays['entities'][str(code)]
        entities.append(dict(canonical_id=f'cl:isapre:{code}', official_code=code, legal_name=row[1],
            commercial_name=overlay['commercial_name'], aliases=sorted(set([row[1]] + overlay['aliases'])),
            isapre_type=overlay['type'], type_source_url=overlays['type_source_url'],
            ambiguous_group_name=overlay.get('ambiguous_group_name', False),
            website=row[2], validity={'status': 'listed_in_official_register', 'observed_at': fetched_at,
                                      'effective_from': None, 'effective_until': None},
            group_control=None, group_control_reason='No control inferred from a shared brand or registry entry',
            source_url=REGISTRY_URL, source_sha256=sha256,
            legal_name_caveat=('SIS register says Ltda.; SIS resolution IF/1315 (2026-02-03) says Isalud Isapre de Codelco SpA; effective corporate change date requires reconciliation' if code == 63 else None)))
    if len(entities) != 10 or {e['official_code'] for e in entities} != EXPECTED_CODES:
        raise ValueError('Official register missing or duplicating an Isapre; preserve previous inventory')
    return entities


def load_universe():
    path = project_root() / 'data/entities/isapres.json'
    return json.loads(path.read_text())['entities'] if path.exists() else []


def resolve_isapres(text, *, codes=(), universe=None):
    """Explicit names only. A parent, generic fundación or essential adjective is not an insurer."""
    text = fold(text); result = []
    for entity in (load_universe() if universe is None else universe):
        explicit = entity['official_code'] in codes
        aliases = entity['aliases']
        if entity.get('ambiguous_group_name'):
            aliases = ['isapre banmedica', 'banmedica isapre']
        if explicit or any(re.search(r'(?<!\w)' + re.escape(fold(alias)) + r'(?!\w)', text) for alias in aliases):
            result.append(entity)
    return result


def entity_refs(text, *, codes=()):
    return [EntityRef(name=e['commercial_name'], institution_type='ISAPRE',
                      subsector=e['isapre_type'], canonical_id=e['canonical_id'])
            for e in resolve_isapres(text, codes=codes)]


def release_metadata(html):
    text = fold(' '.join(read(html).text))
    pub = re.search(r'fecha de publicacion:\s*(\d{1,2}) de (\w+) de (20\d{2})', text)
    period = re.search(r'actualizada? a (\w+) (20\d{2})', text)
    return {'published_at': date(int(pub[3]), MONTHS.index(pub[2])+1, int(pub[1])).isoformat() if pub else None,
            'reporting_period': f'{period[2]}-{MONTHS.index(period[1])+1:02d}' if period else None}


def parse_portfolio(content, *, source_url, sha256, published_at, fetched_at, universe=None):
    """Fail closed on per-month code/metric/total mismatches. No hardcoded latest month."""
    universe = load_universe() if universe is None else universe
    by_code = {e['official_code']: e for e in universe}; observations = []; coverage = []
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        for sheet in book:
            if fold(sheet.title) not in MONTHS: continue
            title = fold(sheet['A4'].value)
            match = re.fullmatch(r'(\w+) (20\d{2})', title)
            if not match or match[1] != fold(sheet.title): raise ValueError('Unknown portfolio reporting period')
            period = f'{match[2]}-{MONTHS.index(match[1])+1:02d}'
            headers = [fold(sheet.cell(6, c).value) for c in (3, 4, 5)]
            if headers != ['n° cotizantes (1)', 'n° cargas (2)', 'n° beneficiarios (1) + (2)']:
                raise ValueError('Unknown portfolio metric headers')
            rows = list(sheet.iter_rows(min_row=8, max_row=24, max_col=5, values_only=True))
            members = [r for r in rows if isinstance(r[0], int) and r[1]]
            # Official annual workbooks include empty future-month templates.
            if members and all(all(v is None for v in r[2:5]) for r in members): continue
            totals = [r for r in rows if r[1] == 'Total Sistema']
            codes = [r[0] for r in members]
            if set(codes) != EXPECTED_CODES or len(codes) != 10 or len(totals) != 1:
                raise ValueError('Missing or duplicate official portfolio codes')
            for r in members + totals:
                if any(not isinstance(n, int) or isinstance(n, bool) or n < 0 for n in r[2:5]):
                    raise ValueError('Non-reconciled portfolio count')
                if r[2] + r[3] != r[4]: raise ValueError('Portfolio counts do not reconcile')
            if any(sum(r[c] for r in members) != totals[0][c] for c in (2, 3, 4)):
                raise ValueError('Official portfolio total does not reconcile')
            denominators = {t: sum(r[4] for r in members if by_code[r[0]]['isapre_type'] == t) for t in ('open', 'closed')}
            for r in members:
                e = by_code[r[0]]; cohort = e['isapre_type']; denominator = denominators[cohort]
                for metric, value, unit in [('beneficiarios', r[4], 'persons'), ('cotizantes', r[2], 'persons'),
                                             ('cargas', r[3], 'persons'), ('beneficiary_share', r[4]/denominator*100 if denominator else None, 'percent')]:
                    if value is None: continue
                    observations.append(dict(entity_id=e['canonical_id'], official_code=r[0], metric_id=metric,
                        value=value, unit=unit, period=period, population=cohort, denominator=denominator,
                        published_at=published_at, source_url=source_url, source_sha256=sha256, fetched_at=fetched_at,
                        sheet=sheet.title, source_row=next(i for i,x in enumerate(rows,8) if x[0] == r[0])))
            coverage.append({'period': period, 'codes': sorted(codes), 'total_beneficiaries': totals[0][4]})
        if not observations: raise ValueError('No validated portfolio months')
    finally: book.close()
    return {'observations': observations, 'coverage': coverage}


def portfolio_benchmark(observations, *, period, cohort):
    """Only compare the same month and open/closed population."""
    return sorted([o for o in observations if o['period'] == period and o['population'] == cohort
                   and o['metric_id'] == 'beneficiary_share'], key=lambda o: o['value'], reverse=True)
