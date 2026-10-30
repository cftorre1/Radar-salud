"""Capture and parse of the SIS 'Estadísticas Financieras del Sistema Isapre' workbook (FEFI) for the F2 ficha.

Same contract as ``stats_capture`` (F4/F6), whose fetcher it reuses: HTTPS on the SIS allowlist only (redirects
included), robots.txt honoured, 25 MB cap, 30 s per request and a deadline for the whole step. Each listed workbook is
downloaded once; only its SHA-256, size, format, sheet names and the values read by the **reviewed template** are
stored in ``data/stats_fichas/captures/fefi.json``, never the binary.

Nothing is read by guesswork (``figure_templates``): the workbook is read only through the structure template in
``config/stats_fichas_v1.json`` → ``families.F2.parser_template``. Per table it pins one sheet, one header row, one
column per (item, period), one row per Isapre and the 'Total sistema' row, each by a pattern that must match exactly
once on normalized text; the unit is anchored the same way; the Isapre rows must add up to the declared total. A
template with a ``<<…>>`` placeholder or without ``reviewed_by``/``reviewed_at`` is not used (status
``template_placeholder`` / ``template_unreviewed``: HOLD). Only a person marks the review, after reading the real
sheet names and header rows that the CI test ``tests/test_stats_ficha_f2_real_xls.py`` prints as annotations.

A captured workbook whose recorded ``template_sha256`` or ``checker_version`` differs from the current one is
re-checked once: the bytes must hash to the recorded SHA-256 (else ``hash_mismatch``: HOLD) and the previous parse is
kept as ``previous_parse``.
"""
from __future__ import annotations

import hashlib
import io
import re
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

from . import figure_templates as ft
from . import stats_capture as cap

SCHEMA = "stats_capture_v1"
KIND = "fefi"
STORE = cap.STORE / "fefi.json"
REGISTRY = Path("config/stats_fichas_v1.json")
CHECKER_VERSION = 1
MAX_SHEETS, MAX_ROWS, MAX_COLS = 60, 800, 100
HEADER_SCAN_ROWS = 80
TITLE_KEY = "estadisticas financieras del sistema isapre"
FORMATS = ("xls", "xlsx")
# The only items a template may map (closed list, so a ficha topic can never be invented from a column name).
ITEMS = ("resultado_ejercicio", "ingresos_operacionales", "costo_operacional", "ingresos_cotizaciones",
         "costo_prestaciones", "costo_sil", "gav", "patrimonio")
PERIODS = ("t", "t_1")


class XlsReaderUnavailable(RuntimeError):
    pass


# ---------------------------------------------------------------------------------------------
# Releases listed in the SIS history


def fefi_releases(root: Path) -> list[dict]:
    """'Estadísticas Financieras del Sistema Isapre a marzo 2026' with exactly one .xls/.xlsx document."""
    out = {}
    for sig in cap._signals(root):
        title = str(sig.get("title") or "")
        if TITLE_KEY not in cap._fold(title) or sig.get("source_name") not in (None, "Superintendencia de Salud"):
            continue
        period = cap.title_period(title)
        docs = [u for u in cap._docs(sig) if u.lower().endswith((".xlsx", ".xls"))]
        if not period or len(docs) != 1:
            continue
        out[period] = {"period": period, "title": title, "release_url": sig.get("source_url"),
                       "published_at": str(sig.get("event_date") or "")[:10] or None, "documents": [{"url": docs[0]}]}
    return [out[k] for k in sorted(out)]


# ---------------------------------------------------------------------------------------------
# Workbook reading (xlsx with openpyxl; legacy xls with xlrd when installed)


def _cell(v: Any) -> Any:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    return v


def read_workbook(data: bytes) -> dict:
    """{'format', 'sheets': [{'name', 'rows': [[cell, …], …]}]}, bounded to MAX_SHEETS × MAX_ROWS × MAX_COLS."""
    fmt = cap.file_format(data)
    sheets = []
    if fmt == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets[:MAX_SHEETS]:
            rows = [[_cell(v) for v in r[:MAX_COLS]] for r in ws.iter_rows(min_row=1, max_row=MAX_ROWS, values_only=True)]
            sheets.append({"name": str(ws.title), "rows": rows})
    elif fmt == "xls":
        try:
            import xlrd  # legacy BIFF reader (pyproject dependency); absent → the capture records it and HOLDs
        except ImportError as exc:
            raise XlsReaderUnavailable("xlrd not installed") from exc
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
        for i in range(min(book.nsheets, MAX_SHEETS)):
            sh = book.sheet_by_index(i)
            rows = [[_cell(v) for v in sh.row_values(r)[:MAX_COLS]] for r in range(min(sh.nrows, MAX_ROWS))]
            sheets.append({"name": str(sh.name), "rows": rows})
    else:
        raise ValueError(f"unexpected_format:{fmt}")
    return {"format": fmt, "sheets": sheets}


def col_letter(i: int) -> str:
    """0-based column index → 'A', 'Z', 'AA'."""
    s, n = "", i + 1
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def row_text(row: list) -> str:
    return " | ".join(ft.normalize_text(v) for v in row if v is not None)


def structure_summary(book: dict, *, header_rows: int = 14, max_chars: int = 2600) -> list[dict]:
    """What a reviewer needs to write the template: each sheet's name, size and its first non-empty rows (normalized,
    cells joined by ' | ', each prefixed by its 1-based row number)."""
    out = []
    for sh in book["sheets"]:
        lines, size = [], 0
        for i, r in enumerate(sh["rows"]):
            txt = row_text(r)
            if not txt:
                continue
            line = f"r{i + 1}: {txt}"
            if len(lines) >= header_rows or size + len(line) > max_chars:
                break
            lines.append(line)
            size += len(line) + 1
        out.append({"name": sh["name"], "rows": len(sh["rows"]),
                    "cols": max((len(r) for r in sh["rows"]), default=0), "head": lines})
    return out


def find_cells(book: dict, sheet_rx: str, text_rx: str, *, limit: int = 6) -> list[dict]:
    """Display helper for CI annotations only (never a data path): cells of the sheets whose name matches
    ``sheet_rx`` and whose normalized text matches ``text_rx``, with their A1 reference."""
    out = []
    for sh in book["sheets"]:
        if not re.search(sheet_rx, ft.normalize_text(sh["name"])):
            continue
        for i, r in enumerate(sh["rows"]):
            for j, v in enumerate(r):
                if v is not None and re.search(text_rx, ft.normalize_text(v)):
                    out.append({"sheet": sh["name"], "cell": f"{col_letter(j)}{i + 1}", "text": str(v)[:1500]})
                    if len(out) >= limit:
                        return out
    return out


# ---------------------------------------------------------------------------------------------
# Template application


NUM_TEXT = re.compile(r"[−-]?\d{1,3}(?:\.\d{3})*(?:,\d+)?|[−-]?\d+(?:,\d+)?")


def numeric(v: Any) -> float | None:
    """A numeric cell, or a string in plain Chilean format ('−1.234,5'); anything else (parentheses, text) is None."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v or "").strip()
    if not NUM_TEXT.fullmatch(s):
        return None
    return float(s.replace("−", "-").replace(".", "").replace(",", "."))


def header_labels(rows: list[list], start: int, n: int) -> list[str]:
    """One label per column for the header block rows[start:start+n]: upper header rows carry a merged label to the
    right ('marzo 2026' spanning two columns); the last header row does not."""
    width = max((len(rows[k]) for k in range(start, min(start + n, len(rows)))), default=0)
    parts = [[] for _ in range(width)]
    for k in range(start, min(start + n, len(rows))):
        carry = None
        last = k == start + n - 1
        for j in range(width):
            v = rows[k][j] if j < len(rows[k]) else None
            txt = ft.normalize_text(v) if v is not None else ""
            if txt:
                carry = txt
            elif not last and carry and j > 0:
                txt = carry
            if txt:
                parts[j].append(txt)
    return [" ".join(p) for p in parts]


def apply_table(book: dict, table_id: str, tpl: dict, entities: dict[str, str],
                subs: dict[str, str] | None = None) -> tuple[dict | None, list[str]]:
    """Read one table through its reviewed template. Returns (table, failures)."""
    fail = lambda why: (None, [f"{table_id}:{why}"])  # noqa: E731
    comp = lambda raw: ft.compile_pattern(raw, subs=subs)  # noqa: E731
    sheet_rx, why = comp(tpl.get("sheet"))
    if not sheet_rx:
        return fail(f"sheet_pattern_{why}")
    idx, _, why = ft.match_once(sheet_rx, [(i, ft.normalize_text(s["name"])) for i, s in enumerate(book["sheets"])])
    if idx is None:
        return fail(f"sheet_{why}")
    sheet = book["sheets"][idx]
    rows = sheet["rows"]
    if tpl.get("orientation", "entities_in_rows") == "entities_in_columns":
        width = max((len(r) for r in rows), default=0)
        rows = [[r[j] if j < len(r) else None for r in rows] for j in range(width)]  # transpose: entities become rows
    elif tpl.get("orientation") not in (None, "entities_in_rows"):
        return fail("orientation_invalid")
    hdr_rx, why = comp(tpl.get("header_row"))
    if not hdr_rx:
        return fail(f"header_pattern_{why}")
    h, _, why = ft.match_once(hdr_rx, [(i, row_text(r)) for i, r in enumerate(rows[:HEADER_SCAN_ROWS])])
    if h is None:
        return fail(f"header_{why}")
    n_hdr = int(tpl.get("header_rows") or 1)
    labels = header_labels(rows, h, n_hdr)
    failures: list[str] = []
    columns: dict[str, dict] = {}
    for col_id, spec in (tpl.get("columns") or {}).items():
        if spec.get("item") not in ITEMS or spec.get("period") not in PERIODS or spec.get("sign", 1) not in (1, -1):
            failures.append(f"{table_id}:column_spec_invalid:{col_id}")
            continue
        rx, why = comp(spec.get("pattern"))
        if not rx:
            failures.append(f"{table_id}:column_pattern_{why}:{col_id}")
            continue
        j, _, why = ft.match_once(rx, list(enumerate(labels)))
        if j is None:
            failures.append(f"{table_id}:column_{why}:{col_id}")
            continue
        columns[col_id] = {"index": j, "letter": col_letter(j), "label": labels[j][:120], "item": spec["item"],
                           "period": spec["period"], "sign": spec.get("sign", 1)}
    if not columns:
        failures.append(f"{table_id}:no_columns")
    if failures:
        return None, failures
    body_start = h + n_hdr
    label_cols = int(tpl.get("label_columns") or 1)
    body = [(i, " ".join(ft.normalize_text(v) for v in rows[i][:label_cols] if v is not None)) for i in range(body_start, len(rows))]
    tot_rx, why = comp(tpl.get("total_row"))
    if not tot_rx:
        return fail(f"total_pattern_{why}")
    total_i = next((i for i, lab in body if lab and tot_rx.search(lab)), None)  # first total row after the header
    if total_i is None:
        return fail("total_no_match")
    region = [(i, lab) for i, lab in body if i < total_i]
    found: dict[str, int] = {}
    for code, pattern in entities.items():
        rx, why = comp(pattern)
        if not rx:
            failures.append(f"{table_id}:entity_pattern_{why}:{code}")
            continue
        i, _, why = ft.match_once(rx, [(i, lab) for i, lab in region if lab])
        if i is None:
            failures.append(f"{table_id}:entity_{why}:{code}")
            continue
        if i in found.values():
            failures.append(f"{table_id}:entity_ambiguous_row:{code}")
        found[code] = i
    mapped = {c["index"] for c in columns.values()}
    for i, lab in region:  # a row with figures in a mapped column must be a listed Isapre (no silently dropped row)
        if i not in found.values() and any(numeric(rows[i][j]) is not None for j in mapped if j < len(rows[i])):
            failures.append(f"{table_id}:unlisted_row:r{i + 1}:{lab[:40]}")
    unit = tpl.get("unit") or {}
    unit_rx, why = comp(unit.get("pattern"))
    scale = unit.get("to_mm_clp")
    if not unit_rx:
        failures.append(f"{table_id}:unit_pattern_{why}")
    elif not isinstance(scale, (int, float)) or isinstance(scale, bool) or scale <= 0:
        failures.append(f"{table_id}:unit_scale_invalid")
    else:
        _, um, why = ft.match_once(unit_rx, [(i, row_text(r)) for i, r in enumerate(sheet["rows"])])
        if um is None:
            failures.append(f"{table_id}:unit_{why}")
    if failures:
        return None, failures

    def ref(i: int, j: int) -> str:
        # A1 reference in the sheet as printed (rows/columns swap back for an 'entities_in_columns' table)
        return f"{col_letter(i)}{j + 1}" if tpl.get("orientation") == "entities_in_columns" else f"{col_letter(j)}{i + 1}"

    values: dict[str, dict] = {}
    cells: dict[str, dict] = {}
    sums: dict[str, str] = {}
    tol = float(tpl.get("sum_tolerance") or 0)
    for col_id, c in columns.items():
        j = c["index"]
        vals, refs = {}, {}
        for code, i in [*found.items(), ("total", total_i)]:
            x = numeric(rows[i][j]) if j < len(rows[i]) else None
            if x is None:
                failures.append(f"{table_id}:non_numeric:{col_id}:{code}:{ref(i, j)}")
                continue
            vals[code], refs[code] = x, ref(i, j)
        if failures:
            continue
        parts = sum(v for k, v in vals.items() if k != "total")
        if abs(parts - vals["total"]) > tol:
            failures.append(f"{table_id}:sum_mismatch:{col_id}")
            continue
        sums[col_id] = f"Σ Isapres {parts:g} = total {vals['total']:g} (tolerancia {tol:g})"
        # 'sign' is declared by the reviewer (e.g. -1 for costs printed as negatives); the sum check runs on the printed values
        values[col_id] = {k: v * scale * c["sign"] for k, v in vals.items()}
        cells[col_id] = refs
    if failures:
        return None, failures
    return {"sheet": sheet["name"], "header_row": h + 1, "unit": unit.get("label") or "MM$", "scale": scale,
            "columns": {k: {kk: vv for kk, vv in c.items() if kk != "index"} for k, c in columns.items()},
            "rows": {**{code: i + 1 for code, i in found.items()}, "total": total_i + 1},
            "values": values, "cells": cells, "sum_check": sums}, []


def coverage_note(book: dict, spec: dict | None, subs: dict[str, str] | None = None) -> tuple[dict | None, str]:
    """The note on coverage that names Esencial (FEFI 'Notas Explicativas'), read through its reviewed pattern."""
    if not spec:
        return None, "not_configured"
    srx, why = ft.compile_pattern(spec.get("sheet"), subs=subs)
    prx, why2 = ft.compile_pattern(spec.get("pattern"), subs=subs)
    if not srx or not prx:
        return None, f"pattern_{why if not srx else why2}"
    idx, _, why = ft.match_once(srx, [(i, ft.normalize_text(s["name"])) for i, s in enumerate(book["sheets"])])
    if idx is None:
        return None, f"sheet_{why}"
    sh = book["sheets"][idx]
    units = [((i, j), ft.normalize_text(v)) for i, r in enumerate(sh["rows"]) for j, v in enumerate(r) if v is not None]
    key, _, why = ft.match_once(prx, units)
    if key is None:
        return None, why
    i, j = key
    text = str(sh["rows"][i][j])
    return {"sheet": sh["name"], "cell": f"{col_letter(j)}{i + 1}", "text": text[:1200],
            "text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest()}, "ok"


def parse_fefi_workbook(data: bytes, template: dict | None, period: str | None = None) -> dict:
    """Hash-bound read of the FEFI workbook through the reviewed template; never a heuristic. ``period`` (the release's
    'YYYY-MM') fills the {t_year}/{t1_year}/{t_month} tokens of the template."""
    base = {"checker_version": CHECKER_VERSION, "template_sha256": ft.templates_sha256(template), "period": period}
    subs = ft.period_subs(period)
    try:
        book = read_workbook(data)
    except XlsReaderUnavailable:
        return {**base, "status": "xls_reader_unavailable", "format": "xls"}
    except Exception as exc:
        return {**base, "status": "workbook_unreadable", "error": f"{type(exc).__name__}"[:80]}
    out = {**base, "format": book["format"], "sheet_names": [s["name"][:80] for s in book["sheets"]][:MAX_SHEETS]}
    state = ft.review_state(template)
    if state != "ok":
        return {**out, "status": f"template_{state}"}
    entities = {str(k): v for k, v in (template.get("entities") or {}).items()}
    tables, failures = {}, []
    if not entities:
        failures.append("entities_not_configured")
    for table_id, tpl in (template.get("tables") or {}).items():
        table, errs = apply_table(book, table_id, tpl, entities, subs)
        failures += errs
        if table:
            tables[table_id] = table
    if not template.get("tables"):
        failures.append("tables_not_configured")
    note, why = coverage_note(book, template.get("coverage_note"), subs)
    if template.get("coverage_note") and not note:
        failures.append(f"coverage_note:{why}")
    return {**out, "status": "verified" if not failures else "failed", "failures": failures[:40],
            "tables": tables if not failures else {}, "coverage_note": note}


# ---------------------------------------------------------------------------------------------
# Capture


def load_template(root: Path) -> dict:
    reg = cap._load(Path(root) / REGISTRY, None)
    if not isinstance(reg, dict):
        from . import stats_fichas as sf
        reg = sf.load_registry()
    return ((reg.get("families") or {}).get("F2") or {}).get("parser_template") or {}


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _capture_doc(doc: dict, *, fetch: Callable[[str], bytes], parse: Callable[[bytes], dict]) -> dict:
    try:
        data = fetch(doc["url"])
    except Exception as exc:
        return {**doc, "status": "fetch_failed", "error": f"{type(exc).__name__}:{exc}"[:160],
                "attempts": int(doc.get("attempts") or 0) + 1, "fetched_at": _now()}
    fmt = cap.file_format(data)
    rec = {"url": doc["url"], "sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data), "format": fmt,
           "fetched_at": _now()}
    if fmt not in FORMATS:
        return {**rec, "status": "unexpected_format", "parse": None}
    try:
        return {**rec, "status": "captured", "parse": parse(data)}
    except Exception as exc:
        return {**rec, "status": "parse_failed", "parse": None, "error": f"{type(exc).__name__}:{exc}"[:160],
                "attempts": int(doc.get("attempts") or 0) + 1}


def _recheck_doc(prev: dict, *, fetch: Callable[[str], bytes], parse: Callable[[bytes], dict],
                 stored: Callable[[dict], bytes | None] | None = None) -> dict:
    """Same rule as F6 (PR #112): bytes from ``stored`` or downloaded again; they must hash to the recorded SHA-256
    (else 'hash_mismatch', HOLD, never re-captured automatically); the previous parse is kept."""
    attempts = int((prev.get("recheck") or {}).get("attempts") or 0) + 1
    try:
        data = (stored(prev) if stored else None) or fetch(prev["url"])
    except Exception as exc:
        return {**prev, "recheck": {"status": "fetch_failed", "error": f"{type(exc).__name__}:{exc}"[:160],
                                    "attempts": attempts, "at": _now()}}
    sha = hashlib.sha256(data).hexdigest()
    if sha != prev.get("sha256"):
        return {**prev, "status": "hash_mismatch", "recheck": {"status": "hash_mismatch", "recorded_sha256": prev.get("sha256"),
                                                              "fetched_sha256": sha, "bytes": len(data), "at": _now()}}
    try:
        parsed = parse(data)
    except Exception as exc:
        return {**prev, "recheck": {"status": "parse_failed", "error": f"{type(exc).__name__}:{exc}"[:160],
                                    "attempts": attempts, "at": _now()}}
    rest = {k: v for k, v in prev.items() if k not in ("recheck", "previous_parse")}
    return {**rest, "status": "captured", "parse": parsed, "previous_parse": prev.get("parse"), "rechecked_at": _now()}


def capture_fefi(root: Path, *, fetch: Callable[[str], bytes] | None = None, max_new: int = 2,
                 time_budget_s: float = cap.TIME_BUDGET_S, clock: Callable[[], float] = time.monotonic,
                 stored: Callable[[dict], bytes | None] | None = None, template: dict | None = None) -> dict:
    """Capture the listed FEFI workbooks (newest first, at most ``max_new`` downloads or re-checks per run)."""
    root = Path(root)
    deadline = clock() + time_budget_s
    fetch = fetch or cap.SisFetcher(deadline=deadline, clock=clock)
    template = load_template(root) if template is None else template
    tpl_sha = ft.templates_sha256(template)
    path = root / STORE
    store = cap._load(path, {}) or {}
    known = {r["period"]: r for r in store.get("releases") or []}
    left = max_new
    summary = {"downloads": 0, "rechecks": 0, "time_budget_exhausted": False}

    def has_time() -> bool:
        if clock() >= deadline:
            summary["time_budget_exhausted"] = True
            return False
        return True

    def days_since(stamp: Any) -> float:
        try:
            return (datetime.now(timezone.utc) - datetime.fromisoformat(str(stamp))).total_seconds() / 86400
        except ValueError:
            return float("inf")

    def needs(old: dict | None) -> bool:
        if not old or old.get("status") == "not_fetched":
            return True
        if old.get("status") == "parse_failed":
            return days_since(old.get("fetched_at")) >= cap.RETRY_AFTER_DAYS
        if old.get("status") != "fetch_failed":
            return False
        return int(old.get("attempts") or 0) < cap.MAX_ATTEMPTS or days_since(old.get("fetched_at")) >= cap.RETRY_AFTER_DAYS

    def needs_recheck(old: dict | None) -> bool:
        if not old or old.get("status") != "captured":
            return False
        p = old.get("parse") or {}
        if int(p.get("checker_version") or 0) >= CHECKER_VERSION and p.get("template_sha256") == tpl_sha:
            return False
        last = old.get("recheck") or {}
        if last.get("status") == "fetch_failed" and int(last.get("attempts") or 0) >= cap.MAX_ATTEMPTS:
            return days_since(last.get("at")) >= cap.RETRY_AFTER_DAYS
        if last.get("status") == "parse_failed":
            return days_since(last.get("at")) >= cap.RETRY_AFTER_DAYS
        return True

    for rel in reversed(fefi_releases(root)):
        old = known.get(rel["period"])
        old_doc = (old or {}).get("document")
        if old_doc and old_doc.get("url") != rel["documents"][0]["url"]:
            old_doc = None  # a new file for the same period (e.g. '_v2'): captured; the revision is detected downstream
        meta = {k: rel[k] for k in ("period", "title", "release_url", "published_at")}
        parse = lambda b, p=rel["period"]: parse_fefi_workbook(b, template, p)  # noqa: E731
        if needs(old_doc) and left > 0 and has_time():
            left -= 1
            summary["downloads"] += 1
            known[rel["period"]] = meta | {"document": _capture_doc({**(old_doc or {}), "url": rel["documents"][0]["url"]},
                                                                    fetch=fetch, parse=parse)}
        elif needs_recheck(old_doc) and left > 0 and has_time():
            left -= 1
            summary["rechecks"] += 1
            known[rel["period"]] = meta | {"document": _recheck_doc(old_doc, fetch=fetch, parse=parse, stored=stored)}
        elif old:
            known[rel["period"]] = {**old, **meta}
    payload = {"schema_version": SCHEMA, "kind": KIND, "updated_at": _now(), "template_sha256": tpl_sha,
               "releases": [known[k] for k in sorted(known)]}
    cap._write(path, payload)
    summary["releases"] = len(payload["releases"])
    summary["verified"] = sum(1 for r in payload["releases"] if ((r.get("document") or {}).get("parse") or {}).get("status") == "verified")
    return summary
