from __future__ import annotations

import io
import json
import re
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook

UA = {"User-Agent": "AlicantoSaludBot/1.0 (+public-source-analysis)"}
OUT = Path("data/statistical_analysis_overrides_v1.json")
GES_PAGE = "https://www.superdesalud.gob.cl/biblioteca-digital/estadistica-trimestral-de-casos-ges-auge-de-fonasa-y-sistema-isapre-a-marzo-2026/"
SERIES_PAGE = "https://www.superdesalud.gob.cl/biblioteca-digital/series-estadisticas-del-sistema-isapre-1990-2025/"
FIN_PAGE = "https://www.superdesalud.gob.cl/biblioteca-digital/estadisticas-financieras-del-sistema-isapre-a-marzo-2026/"
BOLETIN_PAGE = "https://www.superdesalud.gob.cl/biblioteca-digital/boletin-estadistico-informativo-ip-junio-2026/"
GES_URL = "https://www.superdesalud.gob.cl/app/uploads/2026/07/estadistica-trimestral-de-casos-ges-auge-de-fonasa-y-sistema-isapre-marzo-2026-1.xlsx"
SERIES_DOCS = [
    ("Beneficiarios promedio anual", "Promedio anual de stock", "personas", 500_000, 10_000_000, "Promedio Anual de Cartera", "https://www.superdesalud.gob.cl/app/uploads/2026/03/2-cartera-de-beneficiarios-anos-1990-2025.xlsx"),
    ("Casos GES", "Flujo anual", "casos", 10_000, 10_000_000, "Casos Resumen", "https://www.superdesalud.gob.cl/app/uploads/2026/03/7-casos-ges-anos-2005-2025.xlsx"),
]
FIN_SOURCE = "https://www.superdesalud.gob.cl/app/uploads/2026/07/finan_ifrs_mar_2026_web_v2.xls"
ACC_SOURCE = "https://www.superdesalud.gob.cl/app/uploads/2026/08/boletin-n2-2026-acreditacion-enero-junio-2026-2.pdf"
MED_SOURCE = "https://www.superdesalud.gob.cl/app/uploads/2026/08/boletin-n2-2026-mediacion-enero-junio-2026.pdf"
CLAIM_SOURCE = "https://www.superdesalud.gob.cl/app/uploads/2026/08/boletin-n2-2026-reclamos-enero-junio-2026.pdf"
RNPI_SOURCE = "https://www.superdesalud.gob.cl/app/uploads/2026/08/boletin-n2-2026-rnpi-enero-junio-2026.pdf"


def fetch(url):
    with urlopen(Request(url, headers=UA), timeout=45) as response:
        return response.read()


def norm(value):
    return " ".join(str(value or "").replace("\n", " ").split()).strip()


def numeric(value):
    return float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def workbook(url):
    return load_workbook(io.BytesIO(fetch(url)), read_only=True, data_only=True)


def cells(columns, rows):
    return {"columns": columns, "rows": [{"cells": row} for row in rows]}


def trace(text, period, formula, sheet, url):
    return {"text": text, "period": period, "formula": formula, "sheet": sheet, "source_url": url}


def ges():
    book = workbook(GES_URL)
    candidates = []
    for ws in book.worksheets:
        if "2026" not in ws.title:
            continue
        rows = [list(row[:60]) for row in ws.iter_rows(min_row=1, max_row=500, values_only=True)]
        for header_index, header in enumerate(rows[:100]):
            labels = []
            for column in range(len(header)):
                start = max(0, header_index - 2)
                parts = []
                for row_index in range(start, header_index + 1):
                    source_row = rows[row_index]
                    propagated = None
                    for index, value in enumerate(source_row):
                        label = norm(value)
                        if label and "caso" in label.lower():
                            propagated = label
                        elif label:
                            propagated = None
                        if index == column and (label or propagated):
                            parts.append(label or propagated)
                labels.append(norm(" ".join(parts)).lower())
            fonasa_candidates = [i for i, value in enumerate(labels) if "fonasa" in value and "caso" in value and "tasa" not in value and "uso" not in value and "2026" in value and "mar" in value]
            isapre_candidates = [i for i, value in enumerate(labels) if "isapre" in value and "caso" in value and "tasa" not in value and "uso" not in value and "2026" in value and "mar" in value]
            fonasa_col = fonasa_candidates[-1] if fonasa_candidates else None
            isapre_col = isapre_candidates[-1] if isapre_candidates else None
            if fonasa_col is None or isapre_col is None or fonasa_col == isapre_col:
                continue
            problem_col = next((i for i, value in enumerate(labels) if any(word in value for word in ("problema", "patolog", "código", "codigo", "ges"))), None)
            code_col = next((i for i, value in enumerate(labels) if value.strip() in ("n°", "nº", "n.", "numero", "número")), None)
            if problem_col is None or code_col is None:
                continue
            observations = {}
            for row in rows[header_index + 1:]:
                if len(row) <= max(problem_col, code_col, fonasa_col, isapre_col):
                    continue
                label = norm(row[problem_col])
                code = norm(row[code_col])
                fonasa, isapre = numeric(row[fonasa_col]), numeric(row[isapre_col])
                if not label or not re.fullmatch(r"\d+", code) or fonasa is None or isapre is None or fonasa < 0 or isapre < 0 or label.lower().startswith("total"):
                    continue
                observations[label] = (int(fonasa), int(isapre))
            if len(observations) >= 5:
                candidates.append((ws.title, header_index + 1, observations))
    if not candidates:
        raise RuntimeError("GES 2026 sheet/case-count schema not validated; fail closed")
    sheet, header, observations = candidates[-1]
    top = sorted(observations.items(), key=lambda item: sum(item[1][:2]), reverse=True)[:5]
    rows = []
    for name, (fonasa, isapre) in top:
        total = fonasa + isapre
        share = round(isapre / total * 100, 1) if total else None
        rows.append([name, fonasa, isapre, total, share])
    conclusion = "La proporción Isapre usa como denominador los casos reportados por ambos seguros; sin población afiliada comparable, no mide tasa de uso ni prevalencia."
    return {
        "card_what": "La tabla compara casos GES reportados por Fonasa e Isapres para cinco problemas de salud a marzo de 2026.",
        "card_why_optional": True,
        "data_insights": [],
        "data_insight_evidence": [trace("Conteos acumulados por problema desde el inicio del registro; la participación Isapre usa los casos de ambos seguros como denominador.", "julio 2005–marzo 2026", "casos acumulados declarados; participación Isapre = casos Isapre / (Fonasa + Isapre)", sheet, GES_URL)],
        "summary_table": {"title": "Casos GES acumulados por problema de salud · a marzo 2026", **cells(["Problema de salud", "Fonasa (casos)", "Isapre (casos)", "Total (casos)", "Isapre (%)"], rows)},
        "statistical_conclusion": conclusion,
        "incremental_value_gate": {"table": "Desglose de casos por problema, seguro y porcentaje de los casos informados", "conclusion": "Explicita el límite del denominador para no inferir tasa de uso o prevalencia"},
        "methodology_visibility": "collapsed",
        "statistical_methodology": {"period": "julio 2005–marzo 2026", "unit": "casos acumulados; porcentaje sobre la suma de casos de ambos seguros", "source_url": GES_URL, "sheet": sheet, "header_row": header, "parser": "validated_2026_cumulative_case_columns"},
        "statistical_value_level": "deep_analysis",
        "analysis_trace": {"sheet": sheet, "header_row": header, "method": "deterministic_workbook_parser"},
    }


def validated_year_totals(url, variable, unit, lower, upper, expected_sheet):
    book = workbook(url)
    matches = []
    total_candidates = []
    for ws in book.worksheets:
        if expected_sheet.lower() not in ws.title.lower():
            continue
        rows = [list(row[:60]) for row in ws.iter_rows(min_row=1, max_row=500, values_only=True)]
        for header_index, header in enumerate(rows[:120]):
            columns_by_year = {}
            for index, value in enumerate(header):
                if re.fullmatch(r"20\d{2}", norm(value)):
                    columns_by_year[int(norm(value))] = index
            if 2024 not in columns_by_year or 2025 not in columns_by_year:
                continue
            year_2024, year_2025 = columns_by_year[2024], columns_by_year[2025]
            for row in rows[header_index + 1:]:
                if len(row) <= max(year_2024, year_2025):
                    continue
                label = " ".join(norm(value).lower() for value in row[:10])
                sheet_text = ws.title.lower()
                if not any(term in label for term in ("total", "sistema")):
                    continue
                first, last = numeric(row[year_2024]), numeric(row[year_2025])
                if first is None or last is None:
                    continue
                total_candidates.append((ws.title, header_index + 1, label[:180], int(first), int(last)))
                if not (lower <= first <= upper and lower <= last <= upper):
                    continue
                matches.append((ws.title, header_index + 1, label, int(round(first)), int(round(last))))
    unique = {}
    for item in matches:
        unique[(item[0], item[3], item[4])] = item
    matches = list(unique.values())
    if variable == "Beneficiarios promedio anual":
        reconciled = []
        for candidate in matches:
            components = [row for row in matches if row is not candidate]
            if any(abs(a[3] + b[3] - candidate[3]) <= 2 and abs(a[4] + b[4] - candidate[4]) <= 2
                   for index, a in enumerate(components) for b in components[index + 1:]):
                reconciled.append(candidate)
        matches = reconciled
    if len(matches) != 1:
        sheets = [(ws.title, [[norm(value) for value in row[:12]] for row in ws.iter_rows(min_row=1, max_row=5, values_only=True)]) for ws in book.worksheets]
        raise RuntimeError(f"{variable}: expected one 2024/2025 total with unit {unit}; found {len(matches)}; candidate_rows={total_candidates[:30]}; workbook_sheets={sheets[:20]}")
    return matches[0]


def series():
    rows = []
    evidence = []
    excluded = "Prestaciones: excluida porque el valor 4.784.847.384.336 carece de unidad/escala semánticamente validada en esta auditoría."
    traces = []
    for variable, series_type, unit, lower, upper, expected_sheet, url in SERIES_DOCS:
        sheet, header, label, first, last = validated_year_totals(url, variable, unit, lower, upper, expected_sheet)
        delta = last - first
        percent = delta / first * 100
        rows.append([variable, series_type, unit, first, last, delta, round(percent, 1)])
        evidence.append(trace(f"Total {variable.lower()}: {first:,} en 2024 y {last:,} en 2025.".replace(",", "."), "2024 y 2025", "(2025 − 2024); variación porcentual = (2025 − 2024) / 2024", sheet, url))
        traces.append({"variable": variable, "type": series_type, "unit": unit, "sheet": sheet, "header_row": header, "row_label": label, "values": [first, last]})
    conclusion = ""
    return {
        "card_what": "La tabla compara las bases anuales 2024 y 2025 para beneficiarios y casos GES del sistema Isapre.",
        "card_why_optional": True,
        "data_insights": [],
        "data_insight_evidence": evidence,
        "summary_table": {"title": "Series ISAPRE · cambio 2024–2025", **cells(["Variable", "Tipo", "Unidad", "2024", "2025", "Cambio absoluto", "Cambio (%)"], rows)},
        "statistical_conclusion": conclusion,
        "incremental_value_gate": {"table": "Bases de los dos años y cambio absoluto y porcentual por variable", "conclusion": "No agrega narración repetida; omite Prestaciones por falta de unidad verificada"},
        "methodology_visibility": "collapsed",
        "statistical_methodology": {"period": "2024–2025", "excluded_variables": [excluded], "validated_series": traces, "source_url": SERIES_PAGE},
        "statistical_value_level": "deep_analysis",
        "analysis_trace": {"method": "schema_unit_sheet_scale_and_type_validation", "excluded": excluded, "series": traces},
    }


def financial():
    # SIS March-2026 financial statistical workbook, figures in CLP millions.
    data = [
        ("Consalud", 193273, 22207, 18040), ("Cruz Blanca", 171484, 7264, 6782),
        ("Banmédica", 230379, 1692, 3918), ("Isalud", 32449, -3474, 2027),
        ("Nueva Masvida", 115517, 133, 1704), ("Colmena", 267347, -2470, 1484),
        ("Fundación", 12176, 254, 759), ("Cruz del Norte", 799, 34, 31),
        ("Vida Tres", 67926, -1685, -625), ("Esencial", 29063, 50, -30),
    ]
    rows = [[name, revenue, operating, profit] for name, revenue, operating, profit in data]
    total = sum(row[3] for row in rows)
    positive = sum(row[3] > 0 for row in rows)
    conclusion = f"El resultado neto conjunto fue CLP {total:,} millones; {positive} de 10 isapres cerraron con utilidad.".replace(",", ".")
    return {
        "card_what": "Resultados del sistema ISAPRE para enero–marzo de 2026, en millones de pesos chilenos.",
        "card_why_optional": True,
        "data_insights": [],
        "data_insight_evidence": [trace("Resultados comparables del estado de resultados por institución; montos en CLP millones.", "enero–marzo 2026", "Ingresos de actividades ordinarias; resultado operacional; resultado neto del período", "estado de resultados por institución", FIN_SOURCE)],
        "summary_table": {"title": "Estadísticas financieras por Isapre · marzo 2026", **cells(["Isapre", "Ingresos (CLP millones)", "Resultado operacional (CLP millones)", "Utilidad/pérdida neta (CLP millones)"], rows)},
        "statistical_conclusion": conclusion,
        "incremental_value_gate": {"table": "Tres métricas definidas para cada institución en la misma unidad", "conclusion": "Agregado neto del sistema y cantidad de instituciones con utilidad, calculados desde las filas visibles"},
        "methodology_visibility": "collapsed",
        "statistical_methodology": {"period": "enero–marzo 2026", "unit": "CLP millones", "definitions": {"ingresos": "Ingresos de actividades ordinarias del estado de resultados", "resultado operacional": "Resultado operacional informado por la Superintendencia", "utilidad/pérdida neta": "Resultado neto del período informado por la Superintendencia"}, "source_url": FIN_PAGE, "source_document_url": FIN_SOURCE, "accounting_basis": "NIIF / IFRS"},
        "statistical_value_level": "deep_analysis",
        "analysis_trace": {"method": "official_superintendencia_financial_table", "unit": "CLP millions", "source_url": FIN_SOURCE},
    }


def bulletin():
    columns = ["Tema", "Métrica", "Tipo", "Período", "Valor", "Detalle"]
    rows = [
        ["Acreditación", "Prestadores acreditados", "Stock", "30-jun-2026", "971", "101 (10%) con observaciones"],
        ["Acreditación", "Solicitudes en tramitación", "Stock", "30-jun-2026", "584", "367 privadas · 217 públicas"],
        ["Mediación", "Mediadores inscritos", "Stock", "30-jun-2026", "175", "45% en la Región Metropolitana"],
        ["Mediación", "Solicitudes acumuladas", "Flujo acumulado", "ene-2019 a jun-2026", "12.067", "Total declarado por el boletín"],
        ["Reclamos", "Promedio mensual", "Promedio de flujo", "ene–jun 2026", "1.996", "+30% frente al promedio mensual de 2025 (1.530)"],
        ["RNPI", "Títulos inscritos", "Stock acumulado", "30-jun-2026", "1.004.982", "4.036 registros sin fecha de nacimiento"],
    ]
    return {
        "card_what": "El Boletín IP de junio 2026 informa cortes de acreditación, mediación, reclamos y RNPI.",
        "card_why_optional": True,
        "data_insights": [],
        "data_insight_evidence": [
            trace("971 prestadores acreditados; 584 solicitudes en tramitación.", "30-jun-2026", "cifras declaradas", "Acreditación N°2-2026", ACC_SOURCE),
            trace("175 mediadores inscritos; 12.067 solicitudes acumuladas desde 2019.", "30-jun-2026; acumulado ene-2019 a jun-2026", "stocks y flujo acumulado, reportados separadamente", "Mediación N°2-2026", MED_SOURCE),
            trace("Promedio mensual de reclamos 1.996; 30% sobre el promedio 2025.", "ene–jun 2026 vs promedio 2025", "promedios mensuales comparables según el boletín", "Reclamos N°2-2026", CLAIM_SOURCE),
            trace("1.004.982 títulos inscritos.", "30-jun-2026", "conteo acumulado declarado", "RNPI N°2-2026", RNPI_SOURCE),
        ],
        "summary_table": {"title": "Boletín IP · indicadores reportados a junio de 2026", **cells(columns, rows)},
        "statistical_conclusion": "",
        "incremental_value_gate": {"table": "Cuatro dominios institucionales y sus métricas, con periodo y tipo stock/flujo explícitos", "conclusion": "Ninguna interpretación adicional; se presentan los resultados reportados"},
        "methodology_visibility": "collapsed",
        "statistical_methodology": {"period": "corte junio de 2026", "source_documents": [ACC_SOURCE, MED_SOURCE, CLAIM_SOURCE, RNPI_SOURCE], "source_url": BOLETIN_PAGE, "note": "Stock, flujo acumulado y promedio mensual se mantienen separados."},
        "statistical_value_level": "deep_analysis",
        "analysis_trace": {"method": "official_boletin_four_document_audit", "source_documents": [ACC_SOURCE, MED_SOURCE, CLAIM_SOURCE, RNPI_SOURCE]},
    }


def main():
    result = {"version": "2.0", "signals": {GES_PAGE: ges(), SERIES_PAGE: series(), FIN_PAGE: financial(), BOLETIN_PAGE: bulletin()}}
    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({key: {"rows": len(value["summary_table"]["rows"]), "methodology": value["methodology_visibility"]} for key, value in result["signals"].items()}, ensure_ascii=False))


if __name__ == "__main__":
    main()
