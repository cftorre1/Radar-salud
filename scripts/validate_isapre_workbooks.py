"""Controlled, fail-closed validation of the three official 2026 Excel families.

The output is aggregate evidence. It never turns an unverified layout into an insight.
"""
import argparse
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

MONTHS = ("Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio")
URLS = {
    "cartera": "https://www.superdesalud.gob.cl/app/uploads/2026/08/estadistica-mensual-de-cartera-de-beneficiarios-2026.xlsx",
    "suscripciones": "https://www.superdesalud.gob.cl/app/uploads/2026/08/estadistica-mensual-suscripcion-y-desahucios-2026.xlsx",
    "movilidad": "https://www.superdesalud.gob.cl/app/uploads/2026/08/estadistica-mensual-de-movilidad-de-cotizantes-202607.xlsx",
}


def _number(value):
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        raise ValueError("Non-numeric metric in a required row")
    return int(value) if int(value) == value else value


def _count(value):
    number = _number(value)
    if not isinstance(number, int) or number < 0:
        raise ValueError("Required count is fractional or negative")
    return number


def _monthly(book, family):
    result = []
    for index, month in enumerate(MONTHS, 1):
        if month not in book:
            raise ValueError(f"Missing month {month}")
        sheet = book[month]
        if sheet["A4"].value != f"{month.upper()} 2026":
            raise ValueError(f"Unexpected reporting period in {month}")
        headers = [sheet.cell(6, col).value for col in range(1, 6)]
        expected = {2: "N° Cotizantes (1)", 3: "N°\nCargas\n(2)", 4: "N° Beneficiarios (1) + (2)"} if family == "cartera" else {2: "N° Contratos Suscritos", 3: "N°\nDesahucios Voluntarios"}
        if any(headers[col] != label for col,label in expected.items()):
            raise ValueError(f"Unknown {family} schema in {month}")
        rows = [list(r) for r in sheet.iter_rows(min_row=8, max_row=20, max_col=5, values_only=True)]
        members = [r for r in rows if isinstance(r[0], int) and r[1]]
        totals = [r for r in rows if r[1] == "Total Sistema"]
        if len(members) < 7 or len(totals) != 1 or len({r[0] for r in members}) != len(members):
            raise ValueError(f"Missing or duplicate Isapre rows in {month}")
        total = totals[0]
        for col in (2, 3, 4) if family == "cartera" else (2, 3):
            if sum(_count(r[col]) for r in members) != _count(total[col]):
                raise ValueError(f"Total does not reconcile in {month} column {col+1}")
        if family == "cartera":
            if any(_count(r[2]) + _count(r[3]) != _count(r[4]) for r in members):
                raise ValueError(f"Cotizantes plus cargas differs from beneficiaries in {month}")
            metrics = {"cotizantes": _count(total[2]), "cargas": _count(total[3]), "beneficiarios": _count(total[4])}
        else:
            metrics = {"contratos_suscritos": _count(total[2]), "desahucios_voluntarios": _count(total[3])}
        result.append({"period": f"2026-{index:02d}", "sheet": month, "isapres": len(members), "metrics": metrics})
    return result


def _mobility(book):
    if "Nacional" not in book:
        raise ValueError("Missing national mobility sheet")
    sheet = book["Nacional"]
    if not re.fullmatch(r"JULIO 2025 Y JULIO 2026", str(sheet["A3"].value or "")):
        raise ValueError("Unknown mobility period")
    if sheet["C7"].value is not None or sheet["D7"].value != "N° Cotizantes":
        raise ValueError("Unknown national mobility metric")
    labels = {8: "Cotizantes que abandonan el Sistema Isapre", 20: "Cotizantes que ingresan al Sistema Isapre", 32: "Diferencia de Cotizantes"}
    values = {}
    for start, label in labels.items():
        if sheet.cell(start, 2).value != label or sheet.cell(start + 11, 3).value != "Total":
            raise ValueError(f"Unknown mobility block {label}")
        rows = [sheet.cell(row, 4).value for row in range(start, start + 11)]
        values[start] = _number(sheet.cell(start + 11, 4).value)
        if sum(_number(v) for v in rows) != values[start]:
            raise ValueError(f"Mobility age rows do not reconcile: {label}")
    if values[20] - values[8] != values[32]:
        raise ValueError("Mobility net differs from entries minus exits")
    _count(values[8]);_count(values[20])
    return [{"period_start": "2025-07", "period_end": "2026-07", "period_type": "comparison_between_july_cuts", "sheet": "Nacional",
             "metrics": {"salidas_intervalo": values[8], "entradas_intervalo": values[20], "diferencia_intervalo": values[32]}}]


def validate(path: Path, family: str):
    data = path.read_bytes()
    sha = hashlib.sha256(data).hexdigest()
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        series = _mobility(book) if family == "movilidad" else _monthly(book, family)
        return {"family": family, "status": "validated", "source_url": URLS[family], "sha256": sha,
                "bytes": len(data), "schema": "national_comparison_totals_v2" if family == "movilidad" else "monthly_isapre_totals_v2",
                "series": series}
    finally:
        book.close()


def main():
    parser = argparse.ArgumentParser()
    for family in URLS:
        parser.add_argument(f"--{family}", type=Path, required=True)
    parser.add_argument("--output", type=Path, default=Path("data/excel/validated_series.json"))
    args = parser.parse_args()
    output = {"validated_at": datetime.now(timezone.utc).isoformat(), "families": {}}
    for family in URLS:
        try:
            output["families"][family] = validate(getattr(args, family), family)
        except Exception as exc:
            output["families"][family] = {"family": family, "status": "schema_not_validated", "source_url": URLS[family], "error": str(exc)[:180], "series": []}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print({family: item["status"] for family, item in output["families"].items()})
    if any(item["status"] != "validated" for item in output["families"].values()):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
